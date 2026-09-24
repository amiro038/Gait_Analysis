# -*- coding: utf-8 -*-
"""
===============================================================================
 GAIT ANALYSIS -- DICE split-belt treadmill, military load carriage
===============================================================================

 STEP 2 of 2.  Step 1 is detect_gait_events.py (every heel strike and
 toe-off). This script takes those events and computes every gait metric,
 with a 95% confidence interval wherever one can be computed, and writes one
 row per trial into a results workbook.

 HOW TO READ THIS FILE WITH THE WORD DOCUMENT
 -------------------------------------------
 The Word document "Gait_Analysis_Methods.docx" explains every step: the
 equation, why the step is needed, why each choice was made, a figure, and
 the range of values a healthy population shows. Every section of this file
 is headed with the section number it implements, e.g.

        # %% §10  MARGIN OF STABILITY           (Word document §10, Fig. 17)

 so you can read the two side by side. The sections run in the same order as
 the document and in the order the calculation happens:

    §0    settings (every parameter, with the section that uses it)
    §5    preparing a trial
          5.1 events -> steps      5.5 quiet standing and the load
          5.2 belt speed/direction 5.6 system centre of mass
          5.3 boots over the belt  5.7 CoM velocity (markers + force plates)
          5.4 forces into the motion-capture frame
    §6    spatiotemporal            §13  summary statistics, symmetry
    §7    posture, joint motion     §14  long-range correlations (DFA)
    §8    kinetics                  §15  entropy
    §9    CoM work                  §16  trunk: harmonic ratio, regularity
    §10   margin of stability       §17  GEM and foot placement
    §11   foot clearance, trip risk §18  local dynamic stability
    §12   confidence intervals      §19  results tables, figures
          (used by §13-§18)
    RUN   analyse_trial() does §5 -> §19 for one trial; main() does them all

 Sections §2-§4 of the document (equipment, boot meshes, gait events) are
 implemented in build_foot_binding.py and detect_gait_events.py. This file
 borrows the boot geometry from detect_gait_events.py (read_binding,
 boot_sole, toe_hinge, ...), so the events and the metrics see exactly the
 same boot.

 RUNNING IT
 ----------
 In Spyder: set the paths at the top of detect_gait_events.py (they are
 shared), fill in participants.csv, then press F5 here. From a terminal:

    python gait_analysis.py                          # every trial with events
    python gait_analysis.py "D05_C1_Treadmill_1.3mpers 108bpm"
    python gait_analysis.py --tables                 # only rebuild the tables

 WHAT COMES OUT (in OUTPUT_FOLDER)
 ---------------------------------
    all_trials_metrics.xlsx   THE RESULTS: one row per trial, one column per
                              metric; a Key-metrics sheet, one sheet per
                              family, the 95% CIs, the healthy ranges with
                              every trial coloured against them, and a
                              dictionary of every column
    all_trials_metrics.csv    the same values as CSV (+ _ci.csv: 95% limits)
    metric_dictionary.csv     every column: meaning, unit, document section
    all_trials_summary.csv    long form, one row per number (mixed models)
    {trial}_steps.csv         every per-step value, for checking
    {trial}_qa.png            quality-check figure
    {trial}_variability.png   DFA / LDS / entropy figure

 UNITS: metres, seconds, newtons, kilograms, degrees, unless a column name
 says otherwise (_mm, _pct, _bw = per body weight, _tw = per total weight).
===============================================================================
"""

import json
import sys
import traceback
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib
from scipy.interpolate import CubicSpline
from scipy.signal import butter, filtfilt, savgol_filter, sosfiltfilt
from scipy.spatial import ConvexHull, cKDTree

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
# Step 1's script: its paths, its boot geometry and its plate geometry are
# reused here so that events and metrics are computed on the same boot.
import detect_gait_events as dge                                  # noqa: E402

if __name__ == "__main__":
    matplotlib.use("Agg")          # figures go to files, no windows
import matplotlib.pyplot as plt                                   # noqa: E402
from matplotlib.patches import Polygon                            # noqa: E402


# %% ==========================================================================
# §0  SETTINGS                                   (Word document, Appendix A)
# =============================================================================
# Every number that shapes a result is here, with the document section that
# explains it. Change a value here, and the same value is used everywhere.

# --- where things are (the paths are set ONCE, in detect_gait_events.py) ----
FOLDERS = dict(force=dge.FORCE_FOLDER,          # FP_renamed/{trial}.csv
               kinematic=dge.KINEMATIC_FOLDER,  # Theia_csv_outputs/..._metrics.csv
               events=dge.OUTPUT_FOLDER,        # gait_event_outputs/
               binding=dge.BINDING_FILE,        # {participant}_foot_mesh_binding.npz
               plates=dge.PLATE_FILE)           # force_plates_DICE_treadmill.txt
OUTPUT_FOLDER = dge.DATA_FOLDER / "gait_analysis_outputs"
PARTICIPANTS_FILE = HERE / "participants.csv"   # body mass WITHOUT the load
TRIALS = None                  # None = every trial detect_gait_events wrote

# --- sampling (§2) ------------------------------------------------------------
FS = dge.KINEMATIC_RATE        # 100 Hz, Theia
FS_FORCE = dge.GRF_RATE        # 1000 Hz, force plates
STEP = dge.STEP                # force samples per kinematic frame (10)
G = 9.81                       # m/s^2

# --- §5 preparing a trial -------------------------------------------------
WARMUP_S = 60.0                # §5.1 skip the belt ramp and settling
BELT_FLAT = (0.25, 0.55)       # §5.2 foot-flat part of stance (heel rides belt)
FORCE_FILTER_HZ = 50.0         # §5.4 low-pass for force peaks
HANDRAIL_N = 15.0              # §5.4 handrail force that counts as contact
QS_SEARCH_S = 30.0             # §5.5 quiet standing looked for in first 30 s
QS_WINDOW_S = 2.0              # §5.5 ... as the steadiest 2 s window
QS_MAX_CV = 0.02               # §5.5 total vertical force steady within 2%
QS_MIN_SHARE = 0.15            # §5.5 both belts loaded (each >= 15%)
QS_MAX_FOOT_TRAVEL = 0.02      # §5.5 heels still (moved < 20 mm)
LOAD_MIN_KG = 1.0              # §5.6 below 1 kg: no load
LOAD_HEIGHT_M = 0.0            # §5.6 load CoM height above Trunk_Position
LOAD_CENTRED = True            # §5.6 load on the mid-line (left-right)
LOAD_MAX_OFFSET_M = 0.40       # §5.6 load placed further than this: not believed
CROSSOVER_HZ = 0.5             # §5.7 markers below, force plates above
MARKER_SAVGOL = (11, 3)        # §5.7 Savitzky-Golay window, order (markers)
ANTIALIAS_HZ = 40.0            # §5.7 before taking force to 100 Hz

# --- §8-11 per-step metrics ---------------------------------------------------
LOADING_BAND = (0.20, 0.80)    # §8  loading rate between 20% and 80% of F1
COP_FILTER_HZ = 15.0           # §8  CoP and free moment low-pass
WAVE_POINTS = 101              # §7/8 points per time-normalised waveform
MFC_LOCAL_WINDOW = 2           # §11 frames either side a minimum must beat
MFC_SPEED_QUANTILE = 0.75      # §11 MFC must be in the fastest 25% of swing
SWING_TRIM = 0.02              # §11 trim 2% of the swing at each end
MIN_CLEARANCE_MM = 1.0         # §11 floor for the trip-risk division

# --- §12 confidence intervals -------------------------------------------------
N_BOOT = 1000                  # block-bootstrap resamples
N_BOOT_DFA = 200               # simulated series for the DFA interval
CI_LEVEL = 0.95
SEED = 0                       # fixed, so intervals are reproducible

# --- §13-18 variability and stability ------------------------------------------
SERIES_LENGTH = "shortest"     # §12.4 strides per series: "shortest", int, None
LIMB_FOR_TRUNK = "R"           # §16/§18 whose strides segment the trunk signal
DFA_MIN_BOX = 16               # §14 smallest box (Damouras et al. 2010)
DFA_MAX_BOX_FRAC = 1.0 / 9     # §14 largest box = N/9
DFA_N_BOXES = 20               # §14 box sizes, log-spaced
DFA_N_SURROGATE = 20           # §14 shuffled copies (must give alpha ~0.5)
GPH_POWER = 0.5                # §14 GPH bandwidth m = N^0.5
ENT_M = 2                      # §15 template length
ENT_R = 0.2                    # §15 tolerance, x SD
ENT_R_SWEEP = (0.10, 0.15, 0.20, 0.25, 0.30)
MSE_SCALES = 30                # §15 coarse-graining scales (0.3 s)
MSE_MAX_S = 300.0              # §15 seconds of trunk signal for MSE
TRUNK_SOURCES = [("Trunk_Linear_Acceleration", 0),   # §16 (column, number of
                 ("Low_Back_Linear_Acceleration", 0),  #      differentiations
                 ("Trunk_Linear_Velocity", 1),          #      to acceleration)
                 ("Low_Back_Linear_Velocity", 1),
                 ("Low_Back_Position", 2),
                 ("Trunk_Position", 2)]
TRUNK_LOWPASS_HZ = 20.0        # §16
TRUNK_DERIV_WINDOW = 5         # §16 Savitzky-Golay: accurate to 18 Hz
HR_HARMONICS = 20              # §16.1
HR_MIN_SAMPLES = 40            # §16.1 a stride must hold 20 harmonics
REGULARITY_SEARCH = 0.25       # §16.2 peak searched +-25% around expected lag
LDS_N_STRIDES = 150            # §18 strides per window
LDS_SAMPLES_PER_STRIDE = 100   # §18 time normalisation
LDS_TAU = 10                   # §18 delay (normalised samples)
LDS_DE = 5                     # §18 embedding dimension
LDS_WS = 10                    # §18 strides the divergence is followed
LDS_FIT = {"S": (0.0, 0.5), "L": (4.0, 10.0)}   # §18 fit windows, strides
LDS_MAX_WINDOWS = 4            # §18 windows per trial
LDS_N_BOOT = 500               # §18 stride-bootstrap resamples
# §18 state spaces: name -> ([(signal, axis)], dE, scaling). Axis 0/1/2 =
# ML/AP/vertical in the walker's frame. dE None = no delay embedding.
LDS_STATE_SPACES = {
    "trunkVel_ML": ([("Trunk_Linear_Velocity", 0)], LDS_DE, "none"),
    "trunkVel_AP": ([("Trunk_Linear_Velocity", 1)], LDS_DE, "none"),
    "trunkVel_VT": ([("Trunk_Linear_Velocity", 2)], LDS_DE, "none"),
    "trunkVel_3D": ([("Trunk_Linear_Velocity", a) for a in range(3)], 2,
                    "none"),
}
LDS_PRIMARY = "trunkVel_AP"

# --- §19 outputs -----------------------------------------------------------------
RUN_LDS = True
MAKE_FIGURES = True
MAKE_VIDEO = False             # birds-eye MoS video of one typical stride

LIMBS, OTHER, SIDE = ("L", "R"), {"L": "R", "R": "L"}, {"L": "Left",
                                                        "R": "Right"}
HS, TO = "heel_strike", "toe_off"


# %% ==========================================================================
# SMALL TOOLS used throughout
# =============================================================================

def read_export(path):
    """Read a Visual3D metrics export.

    The file has 5 header lines (file name, signal name, type, processing,
    component X/Y/Z or 0..15) and then one row per frame starting with ITEM.
    Each signal is returned as an (F, k) array keyed by its name, so
    columns["Whole_body_COG"] is (F, 3). Trailing empty rows are dropped.
    -> (frame numbers, {name: array}, {name: component labels})
    """
    with open(path, encoding="utf-8-sig", errors="replace") as fh:
        head = [next(fh).rstrip("\r\n").split("\t") for _ in range(5)]
    data = pd.read_csv(path, sep="\t", header=None, skiprows=5,
                       encoding_errors="replace", low_memory=False)
    data = data.apply(pd.to_numeric, errors="coerce").to_numpy(np.float32)
    last = len(data)
    while last and np.isnan(data[last - 1, 1:]).all():
        last -= 1
    data = data[:last]
    groups, comps = {}, {}
    for i, (name, comp) in enumerate(zip(head[1], head[4])):
        if i == 0 or not name.strip():
            continue
        groups.setdefault(name, []).append(i)
        comps.setdefault(name, []).append(comp.strip().upper())
    columns = {n: data[:, idx] for n, idx in groups.items()}
    return data[:, 0].astype(int), columns, comps


def foot_pose(columns, name):
    """A segment's 4x4 pose per frame from its 16 *_Global_4x4 columns.
    The 3x3 rotation is made exactly orthonormal (the export rounds it)."""
    P = columns[name].astype(float).reshape(-1, 4, 4).copy()
    P[:, :3, :3] = dge.orthonormalise(P[:, :3, :3])
    P[:, 3, :] = (0.0, 0.0, 0.0, 1.0)
    P[~np.isfinite(P[:, :3, :]).all(axis=(1, 2))] = np.nan
    return P


def at(x, rows):
    """x (F, ...) at FRACTIONAL rows, by linear interpolation. Events are
    known to a fraction of a frame (§5.1); reading positions at the nearest
    whole frame would throw that precision away."""
    rows = np.asarray(rows, float)
    i = np.clip(np.floor(rows).astype(int), 0, len(x) - 2)
    w = np.clip(rows - i, 0, 1).reshape(rows.shape + (1,) * (x.ndim - 1))
    return x[i] * (1 - w) + x[i + 1] * w


def fill_gaps(x):
    """Fill missing samples by straight lines, column by column, so a filter
    can run across them. Returns the gap mask so the caller can put the
    gaps back afterwards: filled samples are never reported as data."""
    x = np.array(x, float)
    gap = ~np.isfinite(x).all(axis=1) if x.ndim > 1 else ~np.isfinite(x)
    n = np.arange(len(x))
    flat = x.reshape(len(x), -1)
    for j in range(flat.shape[1]):
        ok = np.isfinite(flat[:, j])
        if ok.sum() >= 2:
            flat[~ok, j] = np.interp(n[~ok], n[ok], flat[ok, j])
    return flat.reshape(x.shape), gap


def widen(mask, reach):
    """Grow a True/False mask by `reach` samples on each side (a filter or
    derivative reaches that far, so it is contaminated that far)."""
    return np.convolve(mask, np.ones(2 * reach + 1), "same") > 0


class Trial:
    """Everything known about one trial. load_trial() (§5) fills it in; the
    metric sections read from it and add their columns to trial.steps."""

    fs = FS

    def kin(self, name):
        """A three-component signal (e.g. 'Left_Heel_Position') as (F, 3),
        or None if the export does not have it."""
        if name not in self.columns or len(self.comps[name]) < 3:
            return None
        return self.columns[name][:, :3].astype(float)

    def kin_walker(self, name):
        """The same signal in the WALKER's axes: ML (to the left), AP (along
        the belt's direction of travel), vertical. §5.2 explains why AP/ML
        come from the belt and not from Theia's X and Y."""
        x = self.kin(name)
        if x is None:
            return None
        R = np.column_stack([np.r_[self.lateral, 0], np.r_[self.forward, 0],
                             [0, 0, 1]])
        return x @ R

    def angle(self, name, axis=0):
        """One component of an angle signal in degrees (X = sagittal)."""
        if name not in self.columns:
            return None
        return self.columns[name][:, axis].astype(float)


# %% ==========================================================================
# §5.1  EVENTS -> STEPS                       (Word document §5.1, Fig. 7)
# =============================================================================
# WHAT: turn the list of heel strikes and toe-offs from detect_gait_events.py
#       into one row per heel strike, holding every event that step needs.
# WHY:  every per-step metric is a difference between events: stance is this
#       foot's toe-off minus its heel strike, double support needs the OTHER
#       foot's events in between, and so on. Building the row once, and
#       checking the order of its events once, means every metric uses the
#       same, validated set of events.
# NOTE: events are held as fractional ROWS (0 = first frame) from
#       event_qa.csv's frame_float, so stance and stride times are not
#       rounded to the 10 ms frame grid (a 10 ms rounding is ~1.5% of a
#       stance: larger than many effects of interest).

def read_events(folder, stem, frames):
    """event_qa.csv (sub-frame rows, belt, contact label) if it exists,
    else merged_events.csv (whole frames only)."""
    folder = Path(folder)
    qa = folder / f"{stem}_event_qa.csv"
    if qa.exists():
        ev = pd.read_csv(qa)
        ev["row"] = ev["frame_float"].astype(float)
        ev["grf_sample"] = pd.to_numeric(ev["grf_frame_1000hz"],
                                         errors="coerce")
    else:
        ev = pd.read_csv(dge.find_file(folder, stem, "_merged_events.csv"))
        where = pd.Series(np.arange(len(frames)), index=frames)
        ev["row"] = where.reindex(ev["frame_100hz"]).to_numpy(float)
        ev["grf_sample"] = np.nan
        for c in ("method", "belt", "contact_label"):
            ev[c] = ""
        print(f"  ! no {qa.name}: event times are whole frames and the "
              f"kinetics cannot tell which belt a step was on")
    ev = ev.rename(columns={"support_limb": "limb"})
    # the force sample of an event: exact for GRF events, row x 10 otherwise
    ev["grf_sample"] = ev["grf_sample"].fillna(ev["row"] * STEP)
    return (ev.dropna(subset=["row"]).sort_values("row")
            .reset_index(drop=True))


def build_steps(ev):
    """One row per heel strike (hs), with, as fractional rows:

        to             this foot's toe-off (end of its stance)
        next_hs        this foot's next heel strike (end of its stride)
        contra_to      the OTHER foot's toe-off after hs -> ends the
                       initial double support
        contra_hs      the OTHER foot's heel strike after hs -> starts the
                       terminal double support
        prev_contra_hs the OTHER foot's heel strike before hs -> the step
                       that ENDS at hs started there
    plus where each event came from (GRF / kinematic / ...) and, for GRF
    events, which belt carried it (needed by §8 kinetics).
    """
    def after(rows, x):        # first event in `rows` later than each x
        i = np.searchsorted(rows, x, side="right")
        return np.where(i < len(rows), rows[np.minimum(i, len(rows) - 1)],
                        np.nan)

    def before(rows, x):       # last event in `rows` earlier than each x
        i = np.searchsorted(rows, x, side="left") - 1
        return np.where(i >= 0, rows[np.maximum(i, 0)], np.nan)

    get = {(b, k): ev[(ev["limb"] == b) & (ev["event_type"] == k)]
           for b in LIMBS for k in (HS, TO)}
    rows = {key: g["row"].to_numpy() for key, g in get.items()}
    out = []
    for b in LIMBS:
        o = OTHER[b]
        h = get[(b, HS)]
        s = pd.DataFrame(dict(limb=b, hs=h["row"].to_numpy()))
        s["to"] = after(rows[(b, TO)], s["hs"])
        s["next_hs"] = after(rows[(b, HS)], s["hs"])
        s.loc[s["to"] > s["next_hs"], "to"] = np.nan   # a missed toe-off
        s["contra_to"] = after(rows[(o, TO)], s["hs"])
        s["contra_hs"] = after(rows[(o, HS)], s["hs"])
        s["prev_contra_hs"] = before(rows[(o, HS)], s["hs"])
        s["hs_source"] = h["source"].to_numpy()
        s["hs_belt"] = h["belt"].to_numpy()
        s["hs_sample"] = h["grf_sample"].to_numpy()
        s["label"] = h["contact_label"].to_numpy()
        tos = get[(b, TO)].set_index("row")
        s["to_source"] = tos["source"].reindex(s["to"]).to_numpy()
        s["to_belt"] = tos["belt"].reindex(s["to"]).to_numpy()
        s["to_sample"] = tos["grf_sample"].reindex(s["to"]).to_numpy()
        out.append(s)
    steps = pd.concat(out).sort_values("hs").reset_index(drop=True)
    steps["hs_frame"] = np.round(steps["hs"]).astype(int)
    return steps


# %% ==========================================================================
# §5.2  THE BELT'S SPEED AND DIRECTION OF TRAVEL   (Word document §5.2, Fig. 8)
# =============================================================================
# WHAT: fit a straight line to each stance heel's horizontal position during
#       foot-flat; the velocity of that line is the belt's velocity.
# WHY:  (1) On a treadmill the body barely moves in the lab while the feet
#       sweep backwards. Anything about forward motion (the extrapolated CoM
#       in §10, CoM work in §9) must be expressed relative to the BELT, i.e.
#       as if walking over stationary ground; that needs the belt speed.
#       (2) AP and ML must be along and across the belt. Theia's X/Y need not
#       line up with it, and the boots' long axes toe out (1.3 deg on D05's
#       meshes leaked +-17 mm of step length into step width). A foot on the
#       belt moves exactly along the belt, so its velocity gives the true
#       direction of travel.
# WHY FOOT-FLAT (25-55% of stance): before it the heel is still rolling on,
#       after it the heel lifts; in between the heel is fixed on the belt.

def belt_motion(t, say):
    vel = []
    for s in t.steps.itertuples():
        if not np.isfinite(s.to) or s.to <= s.hs:
            continue
        a = int(s.hs + BELT_FLAT[0] * (s.to - s.hs))
        b = int(s.hs + BELT_FLAT[1] * (s.to - s.hs))
        xy = t.heel[s.limb][a:b, :2]
        if b - a >= 5 and np.isfinite(xy).all():
            # slope of x and of y against time = the heel's velocity
            vel.append(np.polyfit(np.arange(len(xy)) / FS, xy, 1)[0])
    vel = np.array(vel)
    # the heels move BACKWARDS, so forward is minus their median direction
    back = np.median(vel / np.linalg.norm(vel, axis=1, keepdims=True), 0)
    forward = -back / np.linalg.norm(back)
    speed = -vel @ forward
    say(f"  belt: {np.median(speed):.3f} m/s from {len(vel)} stance feet "
        f"(IQR {np.percentile(speed, 25):.3f}-{np.percentile(speed, 75):.3f}), "
        f"travel along ({forward[0]:+.3f}, {forward[1]:+.3f})")
    return forward, float(np.median(speed))


# %% ==========================================================================
# §5.3  BOOTS OVER THE BELT                    (Word document §3, §5.3, Figs 2-3)
# =============================================================================
# WHAT: pose the participant's scanned boot soles on Theia's feet every
#       frame, bend the toe cap at the MTP, and measure them against the
#       belt surface.
# WHY:  the base of support (§10) and foot clearance (§11) are about the
#       real edge and underside of the boot. Joint centres (ankle, "toe")
#       sit centimetres inside it. The boot geometry is the same one
#       detect_gait_events.py used to decide which force-plate events to
#       trust, so the two steps agree about where the boot is.
# KEPT FOR EVERY FRAME (cheap): the lowest sole point's height above the
#       belt and which point it is; the sole's extent along and across the
#       direction of travel (for the base of support). Whole-sole positions
#       are rebuilt only for the frames a swing or a figure needs.

class Boots:
    def __init__(self, trial, verts, summary):
        pose = trial.pose
        self.sole, self.long_axis = {}, {}
        for b in LIMBS:
            # the sole: lowest vertex in every 10 mm cell of the footprint
            sole_v, self.long_axis[b] = dge.boot_sole(verts[b], pose[b])
            # the toe hinge: toe cap turned about the MTP by the toe angle
            angle = trial.angle(f"{SIDE[b]}_Toes_Joint_Angle")
            toes4 = (foot_pose(trial.columns, f"{SIDE[b]}_Toes_Global_4x4")
                     if f"{SIDE[b]}_Toes_Global_4x4" in trial.columns
                     else None)
            toe_pose, toe, _ = dge.toe_hinge(
                pose[b], trial.kin(f"{SIDE[b]}_Toes_Position"), angle, toes4,
                sole_v)
            self.sole[b] = dict(v=sole_v, toe=toe, toe_pose=toe_pose)
        self.pose = pose
        # the belt surface (one slope, one offset per boot), as fitted by the
        # event detection and saved in its summary; refitted if missing
        s = summary.get("belt_surface")
        if s and "walking_direction" in summary:
            fwd = np.asarray(summary["walking_direction"], float)
            self.surface = dict(floor=s["floor_m"], forward=fwd,
                                slope=float(np.tan(np.radians(s["slope_deg"]))))
        else:
            fwd = dge.walking_direction(pose, self.long_axis)
            self.surface = dge.fit_belt_surface(pose, self.sole, fwd)
        fwd2, lat2 = trial.forward, trial.lateral
        n = len(pose["L"])
        self.low_z, self.low_idx, self.ext = {}, {}, {}
        for b in LIMBS:
            self.low_z[b] = np.full(n, np.nan)
            self.low_idx[b] = np.zeros(n, int)
            self.ext[b] = {k: np.full(n, np.nan) for k in
                           ("fwd_max", "fwd_min", "lat_max", "lat_min")}
            for sl, W in dge.sole_chunks(pose[b], self.sole[b]):
                h = self.height(b, W)
                hz = np.where(np.isfinite(h), h, np.inf)
                self.low_idx[b][sl] = hz.argmin(1)
                self.low_z[b][sl] = np.where(np.isfinite(h).any(1),
                                             hz.min(1), np.nan)
                pf, pl = W[..., :2] @ fwd2, W[..., :2] @ lat2
                self.ext[b]["fwd_max"][sl] = pf.max(1)
                self.ext[b]["fwd_min"][sl] = pf.min(1)
                self.ext[b]["lat_max"][sl] = pl.max(1)
                self.ext[b]["lat_min"][sl] = pl.min(1)

    def height(self, b, W):
        """Height of world points W (..., 3) above the pitched belt."""
        return dge.above_belt(W, self.surface, b)

    def world(self, b, rows):
        """All sole points (len(rows), V, 3) at the given integer rows."""
        rows = np.asarray(rows, int)
        P = self.pose[b][rows]
        s = self.sole[b]
        part = dict(s, toe_pose=None if s["toe_pose"] is None
                    else s["toe_pose"][rows])
        return next(dge.sole_chunks(P, part, chunk=len(rows) + 1))[1]


# %% ==========================================================================
# §5.4  FORCES INTO THE MOTION-CAPTURE FRAME       (Word document §5.4, Fig. 9)
# =============================================================================
# WHAT: per belt, remove the unloaded baseline, rotate force, CoP and moment
#       out of the plate's own frame into Theia's, and compute the free
#       moment.
# WHY:  the forces are combined with Theia's CoM (§5.7, §9) and boots (§8),
#       so they must be in the same frame and point the same way. The export
#       gives forces in each plate's frame (DICE: CoP = (-My/Fz, Mx/Fz) -
#       offset); the plate's rotation comes from its measured corners and the
#       lab->Theia registration from the event detection (§4.2).
# SIGN: whether the export is the force ON the plate or ON the body is not
#       documented; the sign is chosen so the vertical force holds the
#       walker up. §5.7 then checks the horizontal axes against the markers.
# BASELINE: an unloaded plate does not read zero, and the reading drifts;
#       it is re-measured every 10 s from the unloaded samples (§4.1).

def read_forces(t, path, plates, A, say):
    head = pd.read_csv(path, nrows=0).columns
    want = [f"{dge.BELT[b]}_{q}_{a}" for b in LIMBS
            for q, axes in (("Force", "XYZ"), ("Moment", "XYZ"), ("COP", "XY"))
            for a in axes]
    rails = [c for c in head if "handrail_Force" in c]
    data = pd.read_csv(path, usecols=[c for c in want + rails if c in head])
    t.grf, t.grf_raw, t.cop, t.free_moment, t.baseline = {}, {}, {}, {}, {}
    sos = butter(4, FORCE_FILTER_HZ, fs=FS_FORCE, output="sos")
    for b in LIMBS:
        belt, pl = dge.BELT[b], plates[b]
        col = lambda q: (data[f"{belt}_{q}"].to_numpy(float)  # noqa: E731
                         if f"{belt}_{q}" in data else None)
        F = np.column_stack([col(f"Force_{a}") for a in "XYZ"])
        # 1. baseline: vertical from the event detection's drift-tracking
        #    baseline, horizontal from the median of the unloaded samples
        base, _ = dge.belt_baseline(F[:, 2])
        loaded = sosfiltfilt(sos, F[:, 2]) - base > dge.FORCE_THRESHOLD_N
        F[:, 2] -= base
        F[:, :2] -= np.median(F[~loaded, :2], axis=0)
        # 2. plate frame -> lab (fitted corners) -> Theia (registration)
        lab = F @ pl["R"].T
        sign = 1.0 if np.mean(lab[loaded, 2]) > 0 else -1.0
        lab *= sign
        theia = np.column_stack([lab[:, :2] @ A[:, :2].T, lab[:, 2]])
        t.grf_raw[b] = theia                            # for integration
        t.grf[b] = sosfiltfilt(sos, theia, axis=0)      # for peaks
        t.baseline[b] = float(np.median(base))
        # 3. CoP into Theia (only meaningful where the foot is loaded)
        cop = np.column_stack([col("COP_X"), col("COP_Y")])
        if dge.COP_FRAME == "plate":
            cop_lab = (np.column_stack([cop, np.zeros(len(cop))])
                       @ pl["R"].T + pl["t"])[:, :2] / 1000
        else:
            cop_lab = cop / 1000
        cop_theia = dge.apply_2d(A, cop_lab)
        cop_theia[~(F[:, 2] > dge.COP_MIN_FORCE_N)] = np.nan
        t.cop[b] = cop_theia
        # 4. free moment: vertical torque NOT explained by the horizontal
        #    forces at the CoP; the CoP must be taken relative to the
        #    plate's moment origin, which is recovered from the data
        M = [col(f"Moment_{a}") for a in "XYZ"]
        if any(m is None for m in M):
            t.free_moment[b] = np.full(len(F), np.nan)
        else:
            Fz = np.where(loaded, F[:, 2] + base, np.nan)
            ox = np.nanmedian(cop[:, 0] + 1000 * M[1] / Fz)
            oy = np.nanmedian(cop[:, 1] - 1000 * M[0] / Fz)
            rx, ry = (cop[:, 0] - ox) / 1000, (cop[:, 1] - oy) / 1000
            fm = M[2] - (rx * F[:, 1] - ry * F[:, 0])
            t.free_moment[b] = np.where(loaded, fm, np.nan)
    t.n_force = len(data)
    # 5. handrails: any real hand force is an external force the GRF misses
    t.handrail = np.zeros(len(data), bool)
    for c in rails:
        x = data[c].to_numpy(float)
        t.handrail |= np.abs(x - np.median(x)) > HANDRAIL_N
    say(f"  forces: {t.n_force / FS_FORCE:.0f} s, baselines "
        f"{t.baseline['L']:.0f} / {t.baseline['R']:.0f} N"
        + (f", handrail touched {t.handrail.mean():.1%} of the time"
           if t.handrail.any() else ""))


# %% ==========================================================================
# §5.5  QUIET STANDING: WEIGHING BODY + LOAD      (Word document §5.5, Fig. 10)
# =============================================================================
# WHAT: find the steadiest 2 s of standing still before the first step and
#       take the mean total vertical force as the weight of body + load.
# WHY:  the load changes between trials, and Theia cannot see it. Standing
#       still, the plates carry exactly the weight of everything the walker
#       has on (Newton: no acceleration). Body mass from participants.csv
#       then gives the load: m_load = W/g - m_body.
# CHECKS: both belts loaded (standing on both feet), heels not moving, force
#       steady to 2%. The mean vertical force while walking must agree with
#       W (over a steady walk the body does not accelerate on average); more
#       than 2% apart means plate drift or a standing that was not still.

def quiet_standing(t, say):
    fz = {b: t.grf[b][:, 2] for b in LIMBS}
    total = fz["L"] + fz["R"]
    first = int(t.events["row"].min() * STEP)           # first event, samples
    end = min(first, int(QS_SEARCH_S * FS_FORCE), len(total))
    win = int(QS_WINDOW_S * FS_FORCE)

    def heels_still(rows):
        """Both heels where they were 2 s earlier (medians of 0.2 s at each
        end, so tracking noise does not count as movement)."""
        k = FS // 5
        for b in LIMBS:
            h = t.heel[b][rows]
            travel = np.linalg.norm(np.nanmedian(h[-k:], 0)
                                    - np.nanmedian(h[:k], 0))
            if not travel < QS_MAX_FOOT_TRAVEL:
                return False
        return True

    best, best_cv = None, np.inf
    for s in range(0, max(end - win, 0), FS_FORCE // 10):   # every 0.1 s
        seg = total[s:s + win]
        share = fz["L"][s:s + win] / seg
        rows = slice(s // STEP, (s + win) // STEP)
        cv = np.std(seg) / np.mean(seg) if np.mean(seg) > 200 else np.inf
        if (cv < min(best_cv, QS_MAX_CV) and share.min() > QS_MIN_SHARE
                and share.max() < 1 - QS_MIN_SHARE and heels_still(rows)):
            best, best_cv = s, cv
    steady = t.steps.loc[t.steps["steady"], "hs"]
    walk = (slice(int(steady.min() * STEP), int(steady.max() * STEP))
            if len(steady) else slice(first, len(total)))
    walking_weight = float(np.mean(total[walk]))

    out = dict(found=best is not None, start_s=np.nan, stop_s=np.nan,
               cv=best_cv, walking_weight_n=walking_weight,
               body_mass_kg=t.body_mass, load_kg=np.nan,
               load_offset_local=np.full(3, np.nan))
    if best is None:
        say("  ! no quiet standing found at the start: the system mass is "
            "taken from the mean vertical GRF while walking, and the load "
            "cannot be placed")
        out.update(weight_n=walking_weight, system_mass_kg=walking_weight / G)
    else:
        sl = slice(best, best + win)
        W = float(np.mean(total[sl]))
        out.update(start_s=best / FS_FORCE, stop_s=(best + win) / FS_FORCE,
                   weight_n=W, system_mass_kg=W / G)
        # the total CoP (both belts, force-weighted) during the standing:
        # §5.6 uses it to place the load
        with np.errstate(invalid="ignore"):
            cop = np.nansum([t.cop[b][sl] * fz[b][sl, None] for b in LIMBS],
                            axis=0) / total[sl, None]
        out["cop_xy"] = np.nanmean(cop, axis=0)
        out["rows"] = (best // STEP, (best + win) // STEP)
        say(f"  quiet standing {out['start_s']:.1f}-{out['stop_s']:.1f} s: "
            f"{W:.0f} N = {W / G:.1f} kg (walking mean "
            f"{walking_weight / W - 1:+.1%})")
        if abs(walking_weight / W - 1) > 0.02:
            say("  ! the walking mean GRF and the standing weight differ by "
                "over 2%: plate drift, or the standing was not still")
    if np.isfinite(t.body_mass):
        out["load_kg"] = out["system_mass_kg"] - t.body_mass
        say(f"  body {t.body_mass:.1f} kg -> load {out['load_kg']:+.1f} kg")
        if out["load_kg"] < -LOAD_MIN_KG:
            say("  ! the plates weigh LESS than the body mass entered in "
                "participants.csv -- check the entry or the plate zero")
    else:
        say("  ! no body mass in participants.csv: the load is unknown and "
            "the CoM is Theia's body CoM")
    return out


# %% ==========================================================================
# §5.6  THE SYSTEM (BODY + LOAD) CENTRE OF MASS    (Word document §5.6, Fig. 11)
# =============================================================================
# WHAT: place the load, carry it with the trunk, and combine it with Theia's
#       body CoM into the CoM of everything the walker moves.
# WHY:  a 20-40 kg load shifts the CoM by several centimetres (backwards and
#       upwards for a pack). Margins of stability and CoM work are about the
#       CoM that must be kept over the feet -- the system's, not the bare
#       body's.
# HOW (standing still the CoP is vertically below the system CoM):
#       x_load = (m * x_CoP - m_body * x_CoM,body) / m_load     (horizontal)
#       height  = Trunk_Position + LOAD_HEIGHT_M   (standing gives no height)
#       left-right = on the mid-line (LOAD_CENTRED): the CoP places the load
#       sideways only to the registration error x m/m_load (~5x), and packs
#       and vests are symmetric.
#       The load's position is stored in a trunk frame (Low_Back -> Neck,
#       x right, y forward) and carried with the trunk through the trial.
#       x_CoM,system = (m_body x_CoM,body + m_load x_load) / m

def trunk_frame(t):
    """(F, 3, 3): columns are the trunk's right, forward and up directions
    (up = Low_Back -> Neck), a right-handed frame."""
    lo, hi = t.kin("Low_Back_Position"), t.kin("Neck_Position")
    up = np.tile([0.0, 0, 1], (t.n, 1))
    if lo is not None and hi is not None:
        u = hi - lo
        ok = np.isfinite(u).all(1)
        up[ok] = u[ok] / np.linalg.norm(u[ok], axis=1, keepdims=True)
    right = -np.r_[t.lateral, 0]
    x = right - (up @ right)[:, None] * up
    x /= np.linalg.norm(x, axis=1, keepdims=True)
    return np.stack([x, np.cross(up, x), up], axis=2)


def system_com(t, say):
    body = t.kin("Whole_body_COG")
    t.com_body = body
    L = t.load
    t.com = body.copy()               # no usable load -> the body's CoM
    t.load_position = None
    if not (L["found"] and np.isfinite(L["load_kg"])
            and L["load_kg"] >= LOAD_MIN_KG):
        return
    trunk = t.kin("Trunk_Position")
    if trunk is None:
        say("  ! no Trunk_Position: the load cannot be carried; CoM is the "
            "body's")
        return
    a, b = L["rows"]
    m, ml, mb = t.mass, L["load_kg"], t.body_mass
    xy = (m * L["cop_xy"] - mb * np.nanmean(body[a:b, :2], 0)) / ml
    z = np.nanmean(trunk[a:b, 2]) + LOAD_HEIGHT_M
    R = trunk_frame(t)
    p = np.r_[xy, z]
    offset = np.nanmedian(np.einsum("fji,fj->fi", R[a:b], p - trunk[a:b]),
                          axis=0)
    if LOAD_CENTRED:
        offset[0] = 0.0
    horizontal = np.hypot(*(p[:2] - np.nanmean(trunk[a:b, :2], 0)))
    if horizontal > LOAD_MAX_OFFSET_M:
        say(f"  ! the load would sit {horizontal * 1000:.0f} mm from the "
            f"trunk; not believed -- placed on the trunk instead")
        offset[:2] = 0.0
    L["load_offset_local"] = offset
    t.load_position = trunk + np.einsum("fij,j->fi", R, offset)
    t.com = (mb * body + ml * t.load_position) / m
    say(f"  load CoM {offset[1] * 1000:+.0f} mm forward of the trunk, "
        + ("centred left-right, " if LOAD_CENTRED else
           f"{offset[0] * 1000:+.0f} mm to the right, ")
        + f"{LOAD_HEIGHT_M * 1000:+.0f} mm up (assumed); system CoM "
        f"{np.nanmedian(np.linalg.norm(t.com - body, axis=1)) * 1000:.0f} mm "
        f"from the body's")


# %% ==========================================================================
# §5.7  CoM VELOCITY: MARKERS + FORCE PLATES       (Word document §5.7, Fig. 12)
# =============================================================================
# WHAT: v = lowpass(d/dt tracked CoM) + highpass(integral of GRF / m)
# WHY:  the margin of stability (§10) divides velocity by w0 (~3.3 /s): 30
#       mm/s of velocity error is 9 mm of margin error, against margins of a
#       few cm. Differentiating a tracked position amplifies noise at the
#       frequencies where the within-stride motion is; integrating the force
#       is exact there (Newton, 1000 Hz, no differentiation) but drifts
#       slowly. Each source is used where it is good, with the SAME filter
#       on both branches so they add up to exactly one (nothing counted
#       twice, nothing lost). Synthetic check: 4-8 mm/s error vs 25 mm/s for
#       the markers alone.
# CHECK: the GRF/m must equal the CoM's acceleration; correlating the two
#       (0.5-5 Hz) per axis tests the force axes, the sign and the mass.

def complementary(v_markers, acc, fs=FS, crossover=None):
    """The complementary filter itself (2nd-order Butterworth, zero lag)."""
    v_f = np.zeros_like(acc)
    v_f[1:] = np.cumsum((acc[:-1] + acc[1:]) / 2, axis=0) / fs   # trapezoid
    v_f -= v_f.mean(0)
    b, a = butter(2, crossover or CROSSOVER_HZ, fs=fs)
    return (filtfilt(b, a, v_markers, axis=0) + v_f
            - filtfilt(b, a, v_f, axis=0))


def fuse_velocity(t, say):
    com, gap = fill_gaps(t.com)
    win, order = MARKER_SAVGOL
    # marker branch: one Savitzky-Golay step smooths and differentiates
    v_mk = savgol_filter(com, win, order, deriv=1, delta=1 / FS, axis=0)
    # force branch: total GRF / system mass. Subtracting the mean removes
    # gravity (the vertical mean IS mg) and plate zero offsets
    total = t.grf_raw["L"] + t.grf_raw["R"]
    acc = (total - total.mean(0)) / t.mass
    aa = butter(4, ANTIALIAS_HZ, fs=FS_FORCE, output="sos")
    acc = sosfiltfilt(aa, acc, axis=0)     # anti-alias before 1000 -> 100 Hz
    rows = np.arange(t.n) * STEP
    a100 = np.full((t.n, 3), np.nan)
    have = rows < len(acc)
    a100[have] = acc[rows[have]]
    rail = np.zeros(t.n, bool)
    rail[have] = t.handrail[rows[have]]
    if (~have).mean() > 0.05:
        say(f"  ! {(~have).mean():.0%} of frames have no force: markers "
            f"alone there")
    a100, _ = fill_gaps(a100)

    # the axes check (Fig. 9)
    band = butter(2, [0.5, 5.0], btype="band", fs=FS, output="sos")
    a_mk = savgol_filter(com, win, order, deriv=2, delta=1 / FS, axis=0)
    t.fusion_check = {}
    for k, name in enumerate(("X", "Y", "Z")):
        r = np.corrcoef(sosfiltfilt(band, a_mk[:, k]),
                        sosfiltfilt(band, a100[:, k]))[0, 1]
        t.fusion_check[name] = float(r)
        if k < 2 and r < -0.5:
            say(f"  ! GRF {name} runs against the CoM acceleration "
                f"(r = {r:.2f}): flipped. Check the force export's axes")
            a100[:, k] *= -1
            for b in LIMBS:
                t.grf[b][:, k] *= -1
                t.grf_raw[b][:, k] *= -1
    say("  GRF vs marker CoM acceleration (0.5-5 Hz): r = "
        + " / ".join(f"{v:.2f}" for v in t.fusion_check.values()))
    if min(t.fusion_check.values()) < 0.7:
        say("  ! weak agreement: the force axes or the mass are off, and "
            "the fused velocity inherits it")

    v = complementary(v_mk, a100)
    bad = widen(gap, win // 2)         # untracked frames stay missing
    v[bad] = np.nan
    t.com_vel = v
    t.com_vel_markers = np.where(bad[:, None], np.nan, v_mk)
    # BELT frame: add the belt speed along the direction of travel (§5.2)
    t.com_vel_belt = v + np.r_[t.forward, 0] * t.belt_speed
    t.handrail_rows = rail
    d = 1000 * np.nanstd(v - v_mk, axis=0)
    say(f"  fused CoM velocity differs from markers alone by "
        f"{d[0]:.0f} / {d[1]:.0f} / {d[2]:.0f} mm/s RMS; belt-frame speed "
        f"{np.nanmean(t.com_vel_belt @ np.r_[t.forward, 0]):.3f} m/s")


# %% ==========================================================================
# §5  PREPARING A TRIAL, in order
# =============================================================================

def load_trial(stem, folders, participant, verbose=True):
    say = print if verbose else (lambda *a, **k: None)
    t = Trial()
    t.stem, t.participant = stem, participant
    say(f"\n{stem}")

    # §5.1 kinematics and events -> steps
    t.frames, t.columns, t.comps = read_export(
        dge.find_file(folders["kinematic"], stem, "_metrics.csv"))
    t.n = len(t.frames)
    t.events = read_events(folders["events"], stem, t.frames)
    say(f"  {t.n / FS:.0f} s of kinematics, {len(t.events)} events "
        f"({(t.events['source'] == 'GRF').mean():.0%} trusted GRF)")
    summary_path = Path(folders["events"]) / f"{stem}_event_summary.json"
    summary = (json.loads(summary_path.read_text())
               if summary_path.exists() else {})
    t.steps = build_steps(t.events)
    first = t.events["row"].min()
    t.steps["steady"] = t.steps["hs"] >= first + WARMUP_S * FS

    # §5.2 belt speed and direction
    t.ankle = {b: t.kin(f"{SIDE[b]}_Ankle_Position") for b in LIMBS}
    t.heel = {b: t.kin(f"{SIDE[b]}_Heel_Position") for b in LIMBS}
    t.forward, t.belt_speed = belt_motion(t, say)
    t.lateral = np.array([-t.forward[1], t.forward[0]])   # to the left

    # §5.3 boots over the belt
    boots, pose_names = dge.read_binding(
        Path(str(folders["binding"]).replace("{participant}",
                                             stem.split("_")[0])))
    t.pose = {b: foot_pose(t.columns, pose_names[b]) for b in LIMBS}
    t.boots = Boots(t, boots, summary)
    n_sole = {b: len(t.boots.sole[b]["v"]) for b in LIMBS}
    say(f"  boots: {n_sole['L']} / {n_sole['R']} sole points; "
        f"belt {t.boots.surface['floor']['L'] * 1000:.1f} / "
        f"{t.boots.surface['floor']['R'] * 1000:.1f} mm, "
        f"{np.degrees(np.arctan(t.boots.surface['slope'])):+.2f} deg")

    # §5.4 forces
    plates = {b: dge.fit_plate(p)
              for b, p in dge.read_plates(folders["plates"]).items()}
    A = np.asarray(summary.get("lab_to_theia", [[1, 0, 0], [0, 1, 0]]), float)
    if "lab_to_theia" not in summary:
        say("  ! no event summary: the plates are taken to be in Theia's "
            "frame")
    read_forces(t, dge.find_file(folders["force"], stem, ".csv"), plates, A,
                say)

    # §5.5 - §5.7 the load, the system CoM, its velocity
    t.body_mass = float(participant.get("body_mass_kg", np.nan))
    t.load = quiet_standing(t, say)
    t.mass = t.load["system_mass_kg"]
    system_com(t, say)
    fuse_velocity(t, say)
    return t


# %% ==========================================================================
# §6  SPATIOTEMPORAL METRICS                   (Word document §6, Figs 7, 13)
# =============================================================================
# One value per heel strike (row of t.steps), from sub-frame event times.
#   stance      to - hs                  swing     next_hs - to
#   stride      next_hs - hs             step      hs - prev_contra_hs
#   cadence     60 / step time (steps/min)
#   double support = (contra_to - hs) + (to - contra_hs)   [initial + terminal]
#   single support = contra_hs - contra_to  (the other foot's swing)
#   percentages are of the STRIDE (the gait cycle), the usual convention
# Supports are only computed when the five events are in their natural
# order; a missed or doubled event would otherwise give a negative time.
#
# Distances use the two HEELS at this heel strike, along and across the
# belt's direction of travel (§5.2):
#   step length = (this heel - other heel) . forward
#   step width  = |(this heel - other heel) . lateral|
# STRIDE LENGTH on a treadmill (Dingwell et al. 2010):
#   L = belt speed x stride time + (landing position now - landing position
#       at the previous heel strike of the same foot)
# i.e. how far the foot travelled over the belt. Two heel-to-heel step
# lengths do NOT add up to it: at the other foot's heel strike this heel has
# already started to lift, which adds 30-50 mm to each step (Fig. 13).

def spatiotemporal(t):
    s = t.steps
    hs, to, nxt = s["hs"], s["to"], s["next_hs"]
    cto, chs, pchs = s["contra_to"], s["contra_hs"], s["prev_contra_hs"]

    def span(a, b, ok=True):
        """(b - a) in seconds, only where b comes after a (and ok)."""
        return ((b - a) / FS).where((b > a) & ok)

    s["stance_s"] = span(hs, to)
    s["swing_s"] = span(to, nxt)
    s["stride_s"] = span(hs, nxt)
    s["step_s"] = span(pchs, hs)
    s["cadence_spm"] = 60 / s["step_s"]
    ordered = (hs < cto) & (cto < chs) & (chs < to) & (to < nxt)
    s["initial_ds_s"] = span(hs, cto, ordered)
    s["terminal_ds_s"] = span(chs, to, ordered)
    s["single_support_s"] = span(cto, chs, ordered)
    s["double_support_s"] = s["initial_ds_s"] + s["terminal_ds_s"]
    for k in ("stance", "swing", "single_support", "double_support"):
        s[f"{k}_pct"] = 100 * s[f"{k}_s"] / s["stride_s"]

    fwd, lat = np.r_[t.forward, 0], np.r_[t.lateral, 0]
    length = np.full(len(s), np.nan)
    width = np.full(len(s), np.nan)
    moved = np.full(len(s), np.nan)
    for b in LIMBS:
        m = (s["limb"] == b).to_numpy()
        rows = s.loc[m, "hs"].to_numpy()
        d = at(t.heel[b], rows) - at(t.heel[OTHER[b]], rows)
        length[m] = d @ fwd
        width[m] = np.abs(d @ lat)
        # how much further forward this heel landed than last time
        e = s.loc[m, "next_hs"].to_numpy()
        ok = np.isfinite(e)
        dm = np.full(m.sum(), np.nan)
        dm[ok] = (at(t.heel[b], e[ok]) - at(t.heel[b], rows[ok])) @ fwd
        moved[m] = dm
    s["step_length_m"] = length
    s["step_width_m"] = width
    s["stride_length_m"] = t.belt_speed * s["stride_s"] + moved
    s["stride_speed_ms"] = s["stride_length_m"] / s["stride_s"]
    s["walk_ratio"] = s["step_length_m"] / s["cadence_spm"]      # §13.3
    # where each step's events came from, for transparency
    s["events_from_kinematics"] = (
        s["hs_source"].isin(["kinematic", "interpolated"])
        | s["to_source"].isin(["kinematic", "interpolated"]))
    s["events_interpolated"] = ((s["hs_source"] == "interpolated")
                                | (s["to_source"] == "interpolated"))
    return s


# %% ==========================================================================
# §7  POSTURE AND JOINT MOTION                        (Word document §7, Fig. 14)
# =============================================================================
# WHY: a load changes posture before it changes timing: the trunk leans
#      forward to bring the system CoM back over the feet, and the knee
#      flexes more in weight acceptance.
# WHAT: per stride (this foot's heel strike to its next), from the first
#      (sagittal) component of Visual3D's angles:
#        trunk lean   mean and range of Thorax_Seg_Angle X (thorax vs lab)
#        pelvis tilt  mean of Pelvis_Seg_Angle X
#        hip/knee/ankle range of motion from <Side>_<Joint>_Joint_Angle X
# Signs are Visual3D's; they are the same in every trial, so differences
# between conditions are valid whichever way "forward" counts.

ANGLES = {                         # output name: signal ({side} = Left/Right)
    "trunk_lean": "Thorax_Seg_Angle",
    "pelvis_tilt": "Pelvis_Seg_Angle",
    "hip": "{side}_Hip_Joint_Angle",
    "knee": "{side}_Knee_Joint_Angle",
    "ankle": "{side}_Ankle_Joint_Angle",
}


def posture(t):
    s = t.steps
    t.angle_waves = {name: [] for name in ANGLES}
    t.angle_waves["limb"], t.angle_waves["hs"] = [], []
    out = {f"{name}_{k}": np.full(len(s), np.nan) for name in ANGLES
           for k in ("mean_deg", "rom_deg")}
    pct = np.linspace(0, 1, WAVE_POINTS)
    for i, st in enumerate(s.itertuples()):
        if not (np.isfinite(st.next_hs) and st.next_hs > st.hs):
            continue
        a, b = int(np.ceil(st.hs)), int(np.floor(st.next_hs))
        for name, signal in ANGLES.items():
            x = t.angle(signal.format(side=SIDE[st.limb]))
            if x is None or b >= len(x) or not np.isfinite(x[a:b + 1]).all():
                t.angle_waves[name].append(np.full(WAVE_POINTS, np.nan))
                continue
            seg = x[a:b + 1]
            out[f"{name}_mean_deg"][i] = seg.mean()
            out[f"{name}_rom_deg"][i] = np.ptp(seg)
            t.angle_waves[name].append(
                np.interp(pct, np.linspace(0, 1, len(seg)), seg))
        t.angle_waves["limb"].append(st.limb)
        t.angle_waves["hs"].append(st.hs)
    for k, v in out.items():
        s[k] = v
    return s


# %% ==========================================================================
# §8  KINETICS and §9 CoM WORK              (Word document §8-9, Figs 15-16)
# =============================================================================
# WHICH STANCES: only those whose heel strike AND toe-off were trusted GRF
#   events of the same belt contact (§4.2): one boot on that belt, the other
#   boot off it, the CoP under the boot. Anything else may carry two feet's
#   force (crossover, shared belt) or be incomplete; it is left missing
#   rather than estimated. The event record names the belt, so a clean
#   crossover step is read from the belt it actually landed on.
#
# §8 per stance (AP = along the direction of travel, positive = forward):
#   F1 / F2        max vertical force in the first / second half of stance
#   trough         min vertical force between F1 and F2
#   loading rate   least-squares slope of vertical force from 20% to 80% of
#                  F1 (a slope over a band is robust to the heel-strike
#                  transient; a two-point slope is not)
#   braking / propulsive impulse   integral of the negative / positive AP
#                  force; they cancel over a stride at steady speed
#   CoP excursion  range of the CoP along / across the belt (15 Hz)
#   free moment    peak |vertical torque| not explained by the horizontal
#                  forces at the CoP (15 Hz)
#   _bw = per body weight (body mass x g); _tw = per total weight (body +
#   load, weighed in §5.5). Load raises _bw values; _tw shows whether the
#   pattern changed beyond carrying more.
#
# §9 individual limbs method (Donelan, Kram & Kuo 2002): each leg's power
#   into the CoM is P = F_leg . v_CoM, with v_CoM the fused SYSTEM velocity
#   in the BELT frame (§5.7), in which the stance foot is still, as over
#   ground. Integrated over the phases of stance:
#     collision  heel strike -> other toe-off        negative work
#     rebound    single support                      positive work
#     preload    single support                      negative work
#     push-off   other heel strike -> toe-off        positive work
#   The step-to-step transition (push-off + collision) is the main
#   mechanical cost of walking and rises with load.

def trusted_stance(st):
    return (st.hs_source == "GRF" and st.to_source == "GRF"
            and isinstance(st.hs_belt, str) and st.hs_belt in LIMBS
            and st.hs_belt == st.to_belt)


def phase_work(power, a, b):
    """(positive, negative) work of power[a:b], in J (1000 Hz samples)."""
    seg = power[max(a, 0):max(b, 0)]
    return (float(np.sum(np.clip(seg, 0, None)) / FS_FORCE),
            float(np.sum(np.clip(seg, None, 0)) / FS_FORCE))


def kinetics(t):
    s = t.steps
    bw = t.body_mass * G if np.isfinite(t.body_mass) else np.nan
    tw = t.load["weight_n"]
    sos = butter(4, COP_FILTER_HZ, fs=FS_FORCE, output="sos")
    fwd = np.r_[t.forward, 0]
    rows = []
    t.waves = {"vgrf": [], "apgrf": [], "limb": [], "hs": []}
    for st in s.itertuples():
        r = {"kinetics_trusted": trusted_stance(st)}
        rows.append(r)
        if not r["kinetics_trusted"]:
            continue
        a, b = int(round(st.hs_sample)), int(round(st.to_sample))
        if b - a < 100 or b > t.n_force:
            continue
        F = t.grf[st.hs_belt][a:b]              # this foot's GRF, N
        ap, vt = F @ fwd, F[:, 2]
        dt = 1 / FS_FORCE
        # --- §8 impulses and peaks ---
        imp = dict(braking_impulse=np.sum(np.clip(ap, None, 0)) * dt,
                   propulsive_impulse=np.sum(np.clip(ap, 0, None)) * dt,
                   vertical_impulse=np.sum(vt) * dt)
        mid = len(vt) // 2
        i1 = int(np.argmax(vt[:mid]))
        i2 = mid + int(np.argmax(vt[mid:]))
        peaks = dict(f1=vt[i1], f2=vt[i2],
                     trough=vt[i1:i2 + 1].min() if i2 > i1 else np.nan,
                     peak_braking=-ap.min(), peak_propulsive=ap.max())
        lo = int(np.argmax(vt[:i1 + 1] >= LOADING_BAND[0] * vt[i1]))
        hi = int(np.argmax(vt[:i1 + 1] >= LOADING_BAND[1] * vt[i1]))
        rate = (np.polyfit(np.arange(lo, hi + 1) * dt, vt[lo:hi + 1], 1)[0]
                if hi - lo >= 3 else np.nan)
        for name, v in (list(imp.items()) + list(peaks.items())
                        + [("loading_rate", rate)]):
            unit = ("_s" if "impulse" in name       # BW*s
                    else "_per_s" if name == "loading_rate" else "")
            r[f"{name}_bw{unit}"] = v / bw
            r[f"{name}_tw{unit}"] = v / tw
        # --- §8 CoP and free moment ---
        cop = t.cop[st.hs_belt][a:b]
        cop = cop[np.isfinite(cop).all(1)]        # only where loaded
        if len(cop) > 30:
            cop = sosfiltfilt(sos, cop, axis=0)
            r["cop_ap_range_mm"] = 1000 * np.ptp(cop @ t.forward)
            r["cop_ml_range_mm"] = 1000 * np.ptp(cop @ t.lateral)
        fm = t.free_moment[st.hs_belt][a:b]
        fm = fm[np.isfinite(fm)]
        if len(fm) > 30:
            r["free_moment_peak_nm"] = float(np.max(np.abs(
                sosfiltfilt(sos, fm))))
        # --- §9 CoM work, individual limbs ---
        v = at(t.com_vel_belt, np.arange(a, b) / STEP)   # velocity at 1 kHz
        if np.isfinite(v).all():
            power = np.sum(F * v, axis=1)                  # W
            k = lambda row: int(round(row * STEP)) - a     # noqa: E731
            r["com_work_pos_j"], r["com_work_neg_j"] = phase_work(
                power, 0, len(power))
            if st.hs < st.contra_to < st.contra_hs < st.to:
                _, r["collision_j"] = phase_work(power, 0, k(st.contra_to))
                r["rebound_j"], r["preload_j"] = phase_work(
                    power, k(st.contra_to), k(st.contra_hs))
                r["push_off_j"], _ = phase_work(power, k(st.contra_hs),
                                                len(power))
            for key in [x for x in list(r) if x.endswith("_j")]:
                r[key.replace("_j", "_j_kg")] = r[key] / t.mass
        # waveforms for the figures and the waveform file
        pct = np.linspace(0, 1, WAVE_POINTS)
        x = np.linspace(0, 1, len(vt))
        t.waves["vgrf"].append(np.interp(pct, x, vt / tw))
        t.waves["apgrf"].append(np.interp(pct, x, ap / tw))
        t.waves["limb"].append(st.limb)
        t.waves["hs"].append(st.hs)
    for k in sorted({k for r in rows for k in r}):
        s[k] = [r.get(k, np.nan) for r in rows]
    s["kinetics_trusted"] = s["kinetics_trusted"].fillna(False).astype(bool)
    return s


# %% ==========================================================================
# §10  MARGIN OF STABILITY                         (Word document §10, Fig. 17)
# =============================================================================
# Hof, Gazendam & Sinke (2005):
#     xCoM = CoM + v / w0,   w0 = sqrt(g / l),   MoS = boundary - xCoM
# If control stopped now, inverted-pendulum dynamics would carry the CoM to
# the extrapolated CoM (xCoM); the margin is how much base of support is
# left beyond it.
#   CoM       the SYSTEM CoM (§5.6)
#   v         fused velocity in the BELT frame (§5.7): the lab-frame CoM
#             hardly moves on a treadmill, so without the belt speed the AP
#             margin would be meaningless
#   l         distance CoM -> stance ankle, frame by frame
#   boundary  the stance BOOT: its outermost sole point across the belt
#             (ML), its most anterior point along it (AP). The ankle joint
#             centre, the usual choice, is centimetres inside the real edge;
#             the ankle-based ML margin is kept (mos_ml_contact_ankle).
# "Outward" = away from the other ankle at heel strike, so the result does
# not depend on which way Theia's axes point.
# Read at heel strike (contact) and as the minimum over SINGLE SUPPORT
# (other toe-off -> other heel strike): only then is the stance boot alone
# the base of support.

def margin_of_stability(t):
    s = t.steps
    cols = ("mos_ml_contact", "mos_ml_min", "mos_ap_contact", "mos_ap_min",
            "mos_ml_min_at_pct", "mos_ml_contact_ankle", "pendulum_m")
    for c in cols:
        s[c] = np.nan
    s["handrail"] = False
    fwd, lat = t.forward, t.lateral
    ext = t.boots.ext
    for i, st in s.iterrows():
        if not (np.isfinite(st.to) and st.to > st.hs):
            continue
        b, o = st.limb, OTHER[st.limb]
        r = np.arange(int(np.ceil(st.hs)), int(np.floor(st.to)) + 1)
        r = r[r < t.n]
        if len(r) < 3:
            continue
        side = np.sign((t.ankle[b][r[0], :2] - t.ankle[o][r[0], :2]) @ lat)
        if not np.isfinite(side) or side == 0:
            continue
        edge = ext[b]["lat_max"][r] if side > 0 else -ext[b]["lat_min"][r]
        length = np.linalg.norm(t.com[r] - t.ankle[b][r], axis=1)
        w0 = np.sqrt(G / length)
        xcom = t.com[r, :2] + t.com_vel_belt[r, :2] / w0[:, None]
        ml = edge - side * (xcom @ lat)
        ap = ext[b]["fwd_max"][r] - xcom @ fwd
        if not np.isfinite(ml).any():
            continue
        single = (r >= st.contra_to) & (r <= st.contra_hs)
        if single.sum() < 3:
            single = np.ones(len(r), bool)
        s.at[i, "mos_ml_contact"] = ml[0]
        s.at[i, "mos_ml_min"] = np.nanmin(ml[single])
        k = np.flatnonzero(single)[np.nanargmin(ml[single])]
        s.at[i, "mos_ml_min_at_pct"] = 100 * k / (len(r) - 1)
        s.at[i, "mos_ap_contact"] = ap[0]
        s.at[i, "mos_ap_min"] = np.nanmin(ap[single])
        s.at[i, "mos_ml_contact_ankle"] = (side * (t.ankle[b][r[0], :2] @ lat)
                                           - side * (xcom[0] @ lat))
        s.at[i, "pendulum_m"] = length[0]
        s.at[i, "handrail"] = bool(t.handrail_rows[r].any())
    leg = t.participant.get("leg_length_m", np.nan)
    if np.isfinite(leg):                  # for comparisons between people
        for c in cols[:4]:
            s[c + "_per_leg"] = s[c] / leg
    return s


# %% ==========================================================================
# §11  FOOT CLEARANCE AND TRIP RISK                (Word document §11, Fig. 18)
# =============================================================================
# MINIMUM FOOT CLEARANCE (MFC): clearance(t) = lowest sole point of the
# swing boot above the pitched belt. The MFC is a LOCAL minimum in mid-swing
# meeting Schulz's (2017) three criteria:
#   (a) lower than the 2 frames either side,
#   (b) the foot is in the fastest 25% of its swing (rules out the dip just
#       after toe-off),
#   (c) the rear of the sole is not lower (the minimum is at the front).
# A swing with no such minimum is a "non-MTC" cycle: missing, not replaced
# by the global minimum (which would be at toe-off or heel strike).
#
# TRIP RISK (Schulz 2017): if the foot caught now, how badly would the body
# be destabilised, relative to how close the foot is to the ground?
#   MoI(t) = max(xCoM_AP - most anterior sole point of EITHER boot, 0)  [mm]
#   TRI    = integral of MoI(t) / clearance(t) [mm/mm] dt   [s]
#   between the MFC point's peak acceleration and peak deceleration (this
#   excludes lift-off and landing, when the foot is meant to be low).
#   The xCoM uses the STANCE leg's pendulum (CoM -> other ankle).

def swing_clearance(t, b, o, rows):
    """MFC, MoI and TRI for one swing of foot b (integer rows)."""
    out = dict(mfc_m=np.nan, mfc_row=np.nan, mfc_vertex=-1, mfc_on_toes=False,
               mfc_minima=0, moi_peak_mm=np.nan, moi_mean_mm=np.nan,
               tri_s=np.nan, tri_window_s=np.nan, tri_peak=np.nan)
    W = t.boots.world(b, rows)               # (frames, sole points, 3)
    H = t.boots.height(b, W)                 # heights above the belt
    if not np.isfinite(H).all():
        return out
    clear = H.min(1)
    low = H.argmin(1)
    speed = np.linalg.norm(np.gradient(W.mean(1), 1 / FS, axis=0), axis=1)
    fast = speed >= np.quantile(speed, MFC_SPEED_QUANTILE)        # (b)
    along = W[..., :2] @ t.forward
    rear = along < along.mean(1, keepdims=True)
    rear_z = np.where(rear, H, np.inf).min(1)                       # (c)
    w = MFC_LOCAL_WINDOW
    cand = []
    for i in range(w, len(clear) - w):                              # (a)
        before, after = clear[i - w:i], clear[i + 1:i + w + 1]
        if (np.all(clear[i] <= before) and np.all(clear[i] <= after)
                and np.any(clear[i] < before) and np.any(clear[i] < after)
                and fast[i] and rear_z[i] >= clear[i]):
            if not cand or i != cand[-1] + 1:          # a flat minimum once
                cand.append(i)
    out["mfc_minima"] = len(cand)

    # margin of instability through the swing
    anterior = np.maximum(t.boots.ext[b]["fwd_max"][rows],
                          t.boots.ext[o]["fwd_max"][rows])
    length = np.linalg.norm(t.com[rows] - t.ankle[o][rows], axis=1)
    xcom = (t.com[rows, :2] @ t.forward
            + (t.com_vel_belt[rows, :2] @ t.forward) / np.sqrt(G / length))
    moi = 1000 * np.maximum(xcom - anterior, 0)
    out["moi_peak_mm"] = float(np.nanmax(moi))
    out["moi_mean_mm"] = float(np.nanmean(moi))
    if not cand:
        return out

    k = min(cand, key=lambda i: clear[i])          # the lowest qualifying one
    v = int(low[k])
    out.update(mfc_m=float(clear[k]), mfc_row=int(rows[k]), mfc_vertex=v,
               mfc_on_toes=bool(t.boots.sole[b]["toe"][v]))
    risk = moi / np.maximum(1000 * clear, MIN_CLEARANCE_MM)
    p = W[:, v]                                     # the MFC point's path
    sp = np.linalg.norm(np.gradient(p, 1 / FS, axis=0), axis=1)
    acc = np.gradient(sp, 1 / FS)
    i0, i1 = sorted((int(np.argmax(acc)), int(np.argmin(acc))))
    if i1 > i0:
        out["tri_s"] = float(np.sum(risk[i0:i1 + 1]) / FS)
        out["tri_window_s"] = (i1 - i0) / FS
    out["tri_peak"] = float(np.nanmax(risk))
    return out


def foot_clearance(t):
    s = t.steps
    results = []
    for st in s.itertuples():
        if not (np.isfinite(st.to) and np.isfinite(st.next_hs)
                and st.next_hs > st.to):
            results.append({})
            continue
        pad = SWING_TRIM * (st.next_hs - st.to)
        rows = np.arange(int(np.ceil(st.to + pad)),
                         int(np.floor(st.next_hs - pad)) + 1)
        rows = rows[rows < t.n]
        results.append(swing_clearance(t, st.limb, OTHER[st.limb], rows)
                       if len(rows) >= 10 else {})
    for k in ("mfc_m", "mfc_row", "mfc_vertex", "mfc_on_toes", "mfc_minima",
              "moi_peak_mm", "moi_mean_mm", "tri_s", "tri_window_s",
              "tri_peak"):
        s[k] = [r.get(k, np.nan) for r in results]
    return s


# %% ==========================================================================
# §12  CONFIDENCE INTERVALS                         (Word document §12, Fig. 19)
# =============================================================================
# Every value comes from ONE walk of a few hundred strides; walk again and
# it changes. A 95% confidence interval says by how much, so a difference
# between conditions can be judged against it. Three methods, because the
# metrics come in three kinds:
#
# 1. CIRCULAR BLOCK BOOTSTRAP (means, SDs, CVs, R2, ...). Strides are not
#    independent: a long stride tends to be followed by another long one.
#    Resampling single strides would destroy that correlation and make the
#    interval too narrow. Instead the series is rebuilt from runs of
#    consecutive strides ("blocks"), drawn at random start points (wrapping
#    round the end), 1000 times; the 2.5th-97.5th percentiles of the
#    recomputed statistic are the interval.
#    BLOCK LENGTH = strides per run. Too short loses the correlation, too
#    long leaves too few distinct blocks. Chosen per series by the rule of
#    Politis & White (2004, corrected 2009) from the series' own
#    autocorrelation: uncorrelated -> 1 (ordinary bootstrap), persistent ->
#    longer. Validated: 93% coverage of a 95% interval on a persistent
#    series, against 66% resampling single strides.
# 2. PARAMETRIC BOOTSTRAP for DFA alpha (§14): blocks would cut exactly the
#    long-range correlation alpha measures, so series with the estimated
#    alpha are SIMULATED (exact fractional Gaussian noise) and DFA is rerun
#    on them.
# 3. STRIDE BOOTSTRAP for the Lyapunov exponent (§18).

def rng_for(tag=""):
    """Random numbers seeded from SEED and a tag: reproducible run to run,
    but different series never share draws."""
    return np.random.default_rng([SEED, sum(map(ord, str(tag)))])


def optimal_block_length(x):
    """Politis & White (2004) / Patton, Politis & White (2009) block length
    for the circular block bootstrap."""
    x = np.asarray(x, float)
    x = x[np.isfinite(x)]
    n = len(x)
    if n < 8 or np.std(x) == 0:
        return 1
    x = x - x.mean()
    kn = max(5, int(np.ceil(np.sqrt(np.log10(n)))))
    m_max = min(int(np.ceil(np.sqrt(n))) + kn, n - 1)
    b_max = int(np.ceil(min(3 * np.sqrt(n), n / 3)))
    acov = np.array([x[:n - k] @ x[k:] / n for k in range(m_max + 1)])
    rho = acov / acov[0]
    # 1. the lag after which kn autocorrelations in a row are negligible
    small = np.abs(rho[1:]) < 2.0 * np.sqrt(np.log10(n) / n)
    m_hat = None
    for m in range(0, m_max - kn + 1):
        if small[m:m + kn].all():
            m_hat = m
            break
    if m_hat is None:
        big = np.flatnonzero(~small)
        m_hat = int(big[-1] + 1) if len(big) else 1
    M = min(2 * max(m_hat, 1), m_max)
    # 2. flat-top lag window estimates of the spectrum at 0 and its slope
    k = np.arange(-M, M + 1)
    tk = np.abs(k) / M
    lam = np.where(tk <= 0.5, 1.0, np.where(tk <= 1.0, 2.0 * (1.0 - tk), 0.0))
    R = acov[np.abs(k)]
    Gk = np.sum(lam * np.abs(k) * R)
    g0 = np.sum(lam * R)
    if g0 <= 0 or Gk == 0:
        return 1
    # 3. the optimal block length for the circular block bootstrap
    b = (2.0 * Gk ** 2 / (4.0 / 3.0 * g0 ** 2)) ** (1.0 / 3.0) * n ** (1.0 / 3.0)
    return int(np.clip(np.round(b), 1, b_max))


def block_indices(n, block, n_boot=None, rng=None):
    """(n_boot, n) resampling indices: runs of `block` consecutive strides
    from random starts, wrapping round the end, glued to length n."""
    n_boot = n_boot or N_BOOT
    rng = rng_for("indices") if rng is None else rng
    block = int(max(1, min(block, n)))
    n_blocks = int(np.ceil(n / block))
    starts = rng.integers(0, n, size=(n_boot, n_blocks))
    return ((starts[:, :, None] + np.arange(block)) % n).reshape(
        n_boot, -1)[:, :n]


def percentile_ci(replicates, level=None):
    level = level or CI_LEVEL
    r = np.asarray(replicates, float)
    r = r[np.isfinite(r)]
    if len(r) < 10:
        return np.nan, np.nan
    a = (1 - level) / 2
    return float(np.quantile(r, a)), float(np.quantile(r, 1 - a))


STATISTICS = {
    "mean": lambda s: np.nanmean(s, axis=-1),
    "sd": lambda s: np.nanstd(s, axis=-1, ddof=1),
    "cv": lambda s: 100 * np.nanstd(s, axis=-1, ddof=1)
    / np.abs(np.nanmean(s, axis=-1)),
}


def series_summary(x, statistics=("mean", "sd", "cv"), n_boot=None, tag=""):
    """Mean, SD and CV of a stride series, each with its block-bootstrap
    interval. -> {statistic: dict(value, ci_low, ci_high, n, block)}"""
    x = np.asarray(x, float)
    x = x[np.isfinite(x)]
    if len(x) < 10:
        return {s: dict(value=np.nan, ci_low=np.nan, ci_high=np.nan,
                        n=len(x), block=np.nan) for s in statistics}
    block = optimal_block_length(x)
    samples = x[block_indices(len(x), block, n_boot, rng_for(tag))]
    out = {}
    with np.errstate(invalid="ignore", divide="ignore"):
        for s in statistics:
            lo, hi = percentile_ci(STATISTICS[s](samples))
            out[s] = dict(value=float(STATISTICS[s](x)), ci_low=lo,
                          ci_high=hi, n=len(x), block=block)
    return out


def bootstrap_rows(n, statistic, block, n_boot=None, tag=""):
    """Interval for any statistic of aligned per-stride arrays; statistic
    takes an index array into the strides and returns a number."""
    idx = block_indices(n, block, n_boot, rng_for(tag))
    return percentile_ci([statistic(i) for i in idx])


def fractional_gaussian_noise(n, hurst, rng):
    """Exact fractional Gaussian noise (Davies & Harte 1987): the series
    whose DFA exponent is alpha = hurst."""
    k = np.arange(0, n + 1)
    gamma = 0.5 * (np.abs(k - 1) ** (2 * hurst) - 2 * np.abs(k) ** (2 * hurst)
                   + np.abs(k + 1) ** (2 * hurst))
    eig = np.maximum(np.fft.fft(np.concatenate([gamma, gamma[-2:0:-1]])).real,
                     0.0)
    m = len(eig)
    z = rng.normal(size=m) + 1j * rng.normal(size=m)
    return np.fft.fft(np.sqrt(eig / (2 * m)) * z).real[:n]


def simulate_alpha(n, alpha, rng):
    """A series with DFA exponent alpha (fGn below 1, its sum above 1)."""
    if alpha < 1.0:
        return fractional_gaussian_noise(n, float(np.clip(alpha, 0.02, 0.98)),
                                         rng)
    return np.cumsum(fractional_gaussian_noise(
        n, float(np.clip(alpha - 1.0, 0.02, 0.98)), rng))


def dfa_interval(alpha_hat, n, n_boot=None, tag=""):
    """Parametric bootstrap interval for DFA alpha: simulate series with
    alpha_hat, rerun DFA, and take the BASIC interval
    [2 a - q97.5, 2 a - q2.5], which corrects DFA's bias at this N.
    -> (ci_low, ci_high, se, bias)"""
    if not np.isfinite(alpha_hat):
        return np.nan, np.nan, np.nan, np.nan
    rng = rng_for(tag)
    reps = np.array([dfa_alpha(simulate_alpha(n, alpha_hat, rng))
                     for _ in range(n_boot or N_BOOT_DFA)])
    reps = reps[np.isfinite(reps)]
    if len(reps) < 20:
        return np.nan, np.nan, np.nan, np.nan
    a = (1 - CI_LEVEL) / 2
    q_lo, q_hi = np.quantile(reps, [a, 1 - a])
    return (float(2 * alpha_hat - q_hi), float(2 * alpha_hat - q_lo),
            float(np.std(reps, ddof=1)), float(np.mean(reps) - alpha_hat))


# %% ==========================================================================
# §13  SUMMARY STATISTICS, SYMMETRY, WALK RATIO     (Word document §13)
# =============================================================================
# Every per-step metric of §6-§11 becomes, per trial, over the STEADY steps
# (after the warm-up, §5.1):
#     <metric>        mean of both feet's steps       (with 95% CI)
#     <metric>_sd     SD                               (with 95% CI)
#     <metric>_cv     CV = SD / |mean| x 100            (with 95% CI)
#     <metric>_L/_R   mean of each foot's steps         (with 95% CI)
# SYMMETRY ANGLE (Zifchock et al. 2008), from the left and right means:
#     SA = (45 deg - atan(X_L / X_R)) / 90 deg x 100%    (0% = symmetric,
#     positive = right larger; -200% if above 100%). Unlike the symmetry
#     index it needs no reference limb and cannot blow up near zero.
# WALK RATIO = step length / cadence (computed per step in §6).

STEP_METRICS = [
    "stride_s", "stance_s", "swing_s", "step_s", "cadence_spm",
    "stance_pct", "double_support_pct", "single_support_pct",
    "step_length_m", "step_width_m", "stride_length_m", "stride_speed_ms",
    "walk_ratio",
    "mos_ml_contact", "mos_ml_min", "mos_ap_contact", "mos_ap_min",
    "mfc_m", "moi_peak_mm", "tri_s",
    "braking_impulse_bw_s", "propulsive_impulse_bw_s", "vertical_impulse_bw_s",
    "f1_bw", "trough_bw", "f2_bw", "loading_rate_bw_per_s",
    "f1_tw", "f2_tw", "propulsive_impulse_tw_s", "braking_impulse_tw_s",
    "peak_braking_bw", "peak_propulsive_bw",
    "cop_ap_range_mm", "cop_ml_range_mm", "free_moment_peak_nm",
    "collision_j_kg", "rebound_j_kg", "preload_j_kg", "push_off_j_kg",
    "com_work_pos_j_kg", "com_work_neg_j_kg",
    "trunk_lean_mean_deg", "trunk_lean_rom_deg", "pelvis_tilt_mean_deg",
    "hip_rom_deg", "knee_rom_deg", "ankle_rom_deg",
]
SYMMETRY_METRICS = ["stance_s", "swing_s", "step_length_m", "mfc_m",
                    "mos_ml_contact", "propulsive_impulse_bw_s", "f1_bw",
                    "tri_s"]


def row(metric, value, limb="both", statistic="value", ci=(np.nan, np.nan),
        n=np.nan, block=np.nan, note=""):
    """One number of the long results table."""
    return dict(metric=metric, limb=limb, statistic=statistic,
                value=float(value) if value is not None else np.nan,
                ci_low=ci[0], ci_high=ci[1], n=n, block=block, note=note)


def steady_steps(t, limb=None):
    s = t.steps[t.steps["steady"]]
    if limb is not None:
        s = s[s["limb"] == limb]
    return s.sort_values("hs")


def trial_descriptors(t):
    """What the trial was, and how trustworthy its data are."""
    ev, s, L = t.events, steady_steps(t), t.load
    out = [row("belt_speed_ms", t.belt_speed),
           row("system_mass_kg", t.mass,
               note="weighed in quiet standing" if L["found"]
               else "mean walking GRF (no quiet standing found)"),
           row("load_kg", L["load_kg"]),
           row("quiet_standing_cv", L["cv"]),
           row("events_trusted_grf_pct", 100 * (ev["source"] == "GRF").mean()),
           row("events_zeni_pct", 100 * (ev["source"] == "kinematic").mean()),
           row("events_interpolated_pct",
               100 * (ev["source"] == "interpolated").mean()),
           row("steps_steady", len(s)),
           row("stances_kinetics_pct", 100 * s["kinetics_trusted"].mean()),
           row("swings_without_mfc_pct", 100 * (s["mfc_minima"] == 0).mean()),
           row("mfc_below_belt_pct", 100 * (s["mfc_m"] < 0).mean()),
           row("handrail_steps_pct", 100 * s["handrail"].mean())]
    for axis, r in t.fusion_check.items():
        out.append(row(f"grf_vs_marker_acc_r_{axis}", r))
    for k, v in zip("xyz", L["load_offset_local"]):
        out.append(row(f"load_offset_{k}_m", v,
                       note="trunk frame: x right, y forward, z up"))
    return out


def step_summaries(t):
    out = []
    for limb in ("L", "R", "both"):
        s = steady_steps(t, None if limb == "both" else limb)
        for m in STEP_METRICS:
            if m not in s:
                continue
            res = series_summary(s[m].to_numpy(float), tag=f"{t.stem}{m}{limb}")
            for stat, r in res.items():
                if np.isfinite(r["value"]):
                    out.append(row(m, r["value"], limb, stat,
                                   (r["ci_low"], r["ci_high"]), r["n"],
                                   r["block"]))
    return out


def symmetry_angle(left, right):
    sa = (45.0 - np.degrees(np.arctan2(left, right))) / 90.0 * 100.0
    return float(sa - 200.0 if sa > 100.0 else sa)


def symmetry_rows(t):
    out = []
    for m in SYMMETRY_METRICS:
        if m not in t.steps:
            continue
        mean = {b: steady_steps(t, b)[m].mean() for b in LIMBS}
        if all(np.isfinite(v) and v != 0 for v in mean.values()):
            out.append(row(f"symmetry_angle_{m}",
                           symmetry_angle(mean["L"], mean["R"]),
                           statistic="pct",
                           note="0 = symmetric, positive = right larger"))
    return out


# %% ==========================================================================
# §14  LONG-RANGE CORRELATIONS: DFA                  (Word document §14, Fig. 20)
# =============================================================================
# Does a stride remember strides hundreds of strides earlier? Detrended
# fluctuation analysis (Peng et al. 1994) on one foot's consecutive strides:
#   1. integrate: Y(k) = sum_{i<=k} (x_i - mean)
#   2. cut Y into boxes of n strides, from the start AND from the end
#   3. remove a straight line from each box (DFA-1)
#   4. F(n) = RMS of what is left, over all boxes
#   5. alpha = slope of log F(n) against log n
# alpha < 0.5 anti-persistent (corrected stride to stride: metronome, fixed
# belt speed), 0.5 uncorrelated, 0.5-1 persistent (healthy overground stride
# time 0.75-0.9), > 1 non-stationary.
# Box sizes 16 to N/9 (Damouras et al. 2010: the popular 4 to N/4 inflates
# alpha). Needs at least an octave of boxes: N >= 288 strides.
# Checks reported with every value: 20 shuffled copies must give ~0.5; the
# Geweke-Porter-Hudak periodogram estimate as an independent second opinion.

# series the metronome paces (their alpha measures beat-following)
CUED = {"stride_s", "stance_s", "swing_s", "step_s", "cadence_spm",
        "stance_pct", "swing_pct", "single_support_pct", "double_support_pct"}
# the stride series DFA and entropy are computed on
SERIES = ["stride_s", "stance_s", "swing_s", "double_support_pct",
          "step_length_m", "step_width_m", "stride_length_m",
          "stride_speed_ms", "mfc_m", "mos_ml_contact", "mos_ml_min",
          "mos_ap_contact", "propulsive_impulse_bw_s", "braking_impulse_bw_s",
          "f1_bw", "trunk_lean_mean_deg"]


def dfa(series, min_box=None, max_box=None, n_boxes=None):
    """-> dict(alpha, r2, n, boxes, fluct). The straight-line residual of
    each box is computed in closed form (fast enough to run thousands of
    times for the interval)."""
    x = np.asarray(series, float)
    x = x[np.isfinite(x)]
    n = len(x)
    min_box = DFA_MIN_BOX if min_box is None else min_box
    n_boxes = DFA_N_BOXES if n_boxes is None else n_boxes
    max_box = int(DFA_MAX_BOX_FRAC * n) if max_box is None else int(max_box)
    empty = dict(alpha=np.nan, r2=np.nan, n=n, boxes=np.array([]),
                 fluct=np.array([]))
    if max_box <= min_box or np.std(x) <= 1e-12 * max(abs(x.mean()), 1.0):
        return empty
    y = np.cumsum(x - x.mean())                                        # 1
    boxes = np.unique(np.round(np.logspace(np.log10(min_box),
                                           np.log10(max_box), n_boxes)).astype(int))
    fluct = np.full(len(boxes), np.nan)
    for i, s in enumerate(boxes):
        count = n // s
        if count < 2:
            continue
        seg = np.vstack([y[:count * s].reshape(count, s),              # 2
                         y[n - count * s:].reshape(count, s)])
        tt = np.arange(s) - (s - 1) / 2
        seg = seg - seg.mean(axis=1, keepdims=True)
        slope = seg @ tt / (tt @ tt)                                   # 3
        resid = np.sum(seg ** 2, axis=1) - slope ** 2 * (tt @ tt)
        fluct[i] = np.sqrt(max(resid.sum(), 0.0) / (2 * count * s))     # 4
    ok = np.isfinite(fluct) & (fluct > 0)
    if ok.sum() < 4:
        return empty
    lx, ly = np.log10(boxes[ok]), np.log10(fluct[ok])
    slope, icept = np.polyfit(lx, ly, 1)                                # 5
    resid = ly - (slope * lx + icept)
    r2 = 1 - np.sum(resid ** 2) / np.sum((ly - ly.mean()) ** 2)
    return dict(alpha=float(slope), r2=float(r2), n=n, boxes=boxes[ok],
                fluct=fluct[ok])


def dfa_alpha(series):
    return dfa(series)["alpha"]


def gph_alpha(series):
    """Geweke & Porter-Hudak: alpha = d + 0.5 from the low-frequency slope
    of the periodogram. About twice DFA's scatter: a ballpark check."""
    x = np.asarray(series, float)
    x = x[np.isfinite(x)]
    n = len(x)
    per = np.abs(np.fft.rfft(x - x.mean())) ** 2 / (2 * np.pi * n)
    w = 2 * np.pi * np.arange(len(per)) / n
    j = np.arange(1, min(int(n ** GPH_POWER), len(per) - 1) + 1)
    with np.errstate(divide="ignore", invalid="ignore"):
        reg, resp = np.log(4 * np.sin(w[j] / 2) ** 2), np.log(per[j])
    ok = np.isfinite(reg) & np.isfinite(resp)
    if ok.sum() < 4:
        return np.nan
    return float(0.5 - np.polyfit(reg[ok], resp[ok], 1)[0])


# %% ==========================================================================
# §15  ENTROPY                                       (Word document §15, Fig. 21)
# =============================================================================
# SAMPLE ENTROPY (Richman & Moorman 2000) of each stride series: how often
# does a pattern of m = 2 consecutive strides that matched within r = 0.2 SD
# still match when extended to m + 1?   SampEn = -ln(A / B)
# (B = pairs matching at length m, A = pairs still matching at m + 1,
# self-matches excluded). Low = regular, high = unpredictable. White noise
# is the MOST unpredictable and the least complex, so a single value is not
# a complexity measure; r is swept 0.10-0.30 to show whether conclusions
# depend on it (Yentes & Raffalt 2021).
# MULTISCALE ENTROPY (refined composite, Wu et al. 2014) of the continuous
# trunk acceleration: coarse-grain by averaging tau samples, for every one of
# the tau starting phases, summing the match counts before the log; r stays
# fixed on the original series. The AREA over scales 1-30 is the complexity.

def match_counts(x, m, r):
    """(A, B): template pairs within r (Chebyshev) at lengths m+1 and m,
    counted by a KD-tree (a double loop is far too slow at 30 000 points)."""
    x = np.asarray(x, float)
    x = x[np.isfinite(x)]
    n = len(x)
    if n < m + 2:
        return 0, 0
    k = n - m                      # the same start points at both lengths
    starts = np.arange(k)
    tm = cKDTree(x[starts[:, None] + np.arange(m)])
    tm1 = cKDTree(x[starts[:, None] + np.arange(m + 1)])
    b = int((tm.count_neighbors(tm, r, p=np.inf) - k) // 2)
    a = int((tm1.count_neighbors(tm1, r, p=np.inf) - k) // 2)
    return a, b


def sample_entropy(series, m=None, r=None):
    """SampEn of the z-scored series (so r is in SD units); NaN, not 0, when
    nothing matches."""
    m = ENT_M if m is None else m
    r = ENT_R if r is None else r
    x = np.asarray(series, float)
    x = x[np.isfinite(x)]
    if len(x) < m + 2 or np.std(x) <= 1e-12 * max(abs(x.mean()), 1.0):
        return np.nan
    a, b = match_counts((x - x.mean()) / np.std(x), m, r)
    return float(-np.log(a / b)) if a > 0 and b > 0 else np.nan


def rcmse(series, scales=None, m=None, r=None):
    """Refined composite multiscale entropy curve (scales 1..scales)."""
    scales = scales or MSE_SCALES
    m = ENT_M if m is None else m
    r = ENT_R if r is None else r
    x = np.asarray(series, float)
    x = x[np.isfinite(x)]
    x = (x - x.mean()) / np.std(x)
    out = np.full(scales, np.nan)
    for s in range(1, scales + 1):
        a_sum = b_sum = 0
        for p in range(s):                       # every coarse-graining phase
            cut = x[p:]
            k = len(cut) // s
            if k < m + 2:
                continue
            a, b = match_counts(cut[:k * s].reshape(k, s).mean(axis=1), m, r)
            a_sum, b_sum = a_sum + a, b_sum + b
        if a_sum > 0 and b_sum > 0:
            out[s - 1] = -np.log(a_sum / b_sum)
    return out


def stride_series_metrics(series, tag=""):
    """§14 DFA (+ interval, shuffle check, GPH) and §15 sample entropy (+ r
    sweep) for every column of one foot's consecutive steady strides."""
    rows = []
    for name in series.columns:
        x = pd.to_numeric(series[name], errors="coerce").to_numpy(float)
        ok = np.isfinite(x)
        if ok.mean() < 0.95 or ok.sum() < 50:
            continue                 # a series with gaps is not a series
        x = x[ok]
        if np.std(x) <= 1e-12 * max(abs(x.mean()), 1.0):
            continue
        res = (dfa(x) if DFA_MAX_BOX_FRAC * len(x) >= 2 * DFA_MIN_BOX
               else dict(alpha=np.nan))
        lo, hi, se, bias = dfa_interval(res["alpha"], len(x),
                                        tag=f"{tag}{name}")
        base = dict(series=name, n=len(x), cued=name in CUED)
        if np.isfinite(res["alpha"]):
            rng = rng_for(f"shuffle{tag}{name}")
            sur = [dfa_alpha(rng.permutation(x))
                   for _ in range(DFA_N_SURROGATE)]
            rows.append(dict(base, metric="dfa_alpha", value=res["alpha"],
                             ci_low=lo, ci_high=hi, se=se, bias=bias,
                             r2=res["r2"], gph=gph_alpha(x),
                             shuffled=float(np.nanmean(sur)),
                             boxes=f"{res['boxes'][0]}-{res['boxes'][-1]}"))
        rows.append(dict(base, metric="sample_entropy",
                         value=sample_entropy(x),
                         **{f"r{r:.2f}": sample_entropy(x, r=r)
                            for r in ENT_R_SWEEP}))
    return rows


# %% ==========================================================================
# §16  TRUNK: HARMONIC RATIO, REGULARITY            (Word document §16, Fig. 22)
# =============================================================================
# THE SIGNAL: the first LINEAR trunk signal in the export
# (Trunk_Linear_Velocity on DICE). NOT *_Joint_Acc / Trunk_Joint_Acceleration:
# those are ANGULAR accelerations in deg/s^2. Low-pass 20 Hz, rotated into
# the walker's ML / AP / vertical axes; the time-domain acceleration is a
# short Savitzky-Golay derivative (accurate to 18 Hz). Gaps stay missing.
#
# §16.1 HARMONIC RATIO per stride (right-foot heel strike to heel strike):
# the Fourier transform of exactly one stride, so bin k is harmonic k. AP
# and vertical repeat twice per stride (even harmonics = intrinsic), ML once
# (odd = intrinsic):  HR = sum(intrinsic amplitudes) / sum(the others),
# over harmonics 1-20. From a velocity, acceleration harmonic k is exactly
# k*w times velocity harmonic k, so amplitudes are weighted by k -- exact,
# where a numerical derivative would attenuate the high harmonics.
# iHR (Pasciuto et al. 2017) = intrinsic POWER / total power x 100: bounded
# 0-100%, so it can be averaged; HR is unbounded, so its median is used.
#
# §16.2 REGULARITY (Moe-Nilssen & Helbostad 2004): the unbiased
# autocorrelation of the trunk acceleration, read at its PEAK near one step
# (step regularity) and near one stride (stride regularity); their ratio is
# the symmetry. Peaks, not fixed lags: a 6% drift in stride time makes a
# perfectly regular signal read 0.94 at a fixed lag.

def trunk_signal(t):
    """-> dict(acc = time-domain acceleration (ML, AP, VT), source = the
    filtered signal the harmonic ratio uses, power = differentiations
    needed, name) or None."""
    for name, power in TRUNK_SOURCES:
        raw = t.kin(name)
        if raw is not None:
            break
    else:
        return None
    gap = ~np.isfinite(raw).all(axis=1)
    fill, _ = fill_gaps(raw)
    sos = butter(4, TRUNK_LOWPASS_HZ, fs=FS, output="sos")
    source = sosfiltfilt(sos, fill, axis=0)
    acc = source.copy()
    for _ in range(power):
        acc = savgol_filter(acc, TRUNK_DERIV_WINDOW, 3, deriv=1,
                            delta=1.0 / FS, axis=0)
    axes = np.column_stack([np.r_[t.lateral, 0], np.r_[t.forward, 0],
                            [0, 0, 1]])                     # -> ML, AP, VT
    acc, source = acc @ axes, source @ axes
    acc[widen(gap, power * (TRUNK_DERIV_WINDOW // 2))] = np.nan
    source[gap] = np.nan
    return dict(acc=acc, source=source, power=power, name=name)


def harmonics(segment, n_harmonics=None, power=0):
    """Amplitudes of harmonics 1..20 of one stride, weighted k**power."""
    n_harmonics = n_harmonics or HR_HARMONICS
    seg = np.asarray(segment, float)
    amp = np.abs(np.fft.rfft(seg - seg.mean()))
    if len(amp) <= n_harmonics:
        return None
    k = np.arange(1, n_harmonics + 1)
    return amp[1:n_harmonics + 1] * k.astype(float) ** power


def harmonic_ratio(amp, odd_dominant=False):
    """-> (HR, iHR %). odd_dominant for ML."""
    k = np.arange(1, len(amp) + 1)
    odd, even = amp[k % 2 == 1], amp[k % 2 == 0]
    intrinsic, other = (odd, even) if odd_dominant else (even, odd)
    hr = intrinsic.sum() / other.sum() if other.sum() > 0 else np.nan
    total = np.sum(amp ** 2)
    ihr = 100 * np.sum(intrinsic ** 2) / total if total > 0 else np.nan
    return float(hr), float(ihr)


def unbiased_autocorr(x, max_lag):
    """Autocorrelation divided by the OVERLAP at each lag (not by N), so a
    longer lag is not penalised just for being longer."""
    x = np.asarray(x, float)
    x = x[np.isfinite(x)]
    x = x - x.mean()
    n = len(x)
    ac = np.full(max_lag + 1, np.nan)
    for k in range(min(max_lag, n - 10) + 1):
        ac[k] = x[:n - k] @ x[k:] / (n - k)
    return ac / ac[0]


def regularity(ac, stride_lag):
    """Step and stride regularity at the autocorrelation PEAKS near half a
    stride and a stride, and their ratio (symmetry)."""
    def peak(lag):
        lo = int(np.floor(lag * (1 - REGULARITY_SEARCH)))
        hi = min(int(np.ceil(lag * (1 + REGULARITY_SEARCH))), len(ac) - 1)
        if hi <= lo:
            return np.nan, np.nan
        i = lo + int(np.nanargmax(ac[lo:hi + 1]))
        return float(ac[i]), i
    d1, l1 = peak(stride_lag / 2)
    d2, l2 = peak(stride_lag)
    return dict(step_regularity=d1, stride_regularity=d2,
                symmetry=d1 / d2 if d2 else np.nan, step_lag=l1,
                stride_lag=l2)


def trunk_rows(t, limb, n):
    """§16 (and the §15 multiscale entropy) for one trial."""
    sig = trunk_signal(t)
    if sig is None:
        return [], None
    s = series_for(t, limb, n)
    s = s[np.isfinite(s["next_hs"])]
    strides = list(zip(s["hs"], s["next_hs"]))
    acc, src, power = sig["acc"], sig["source"], sig["power"]
    # §16.1 harmonic ratio, stride by stride
    per = []
    for a, b in strides:
        a, b = int(round(a)), int(round(b))
        if b - a < HR_MIN_SAMPLES or b > len(src):
            continue
        seg = src[a:b]
        if not np.isfinite(seg).all():
            continue
        r = dict(start=a)
        for ax, label, odd in ((1, "AP", False), (2, "VT", False),
                               (0, "ML", True)):
            amp = harmonics(seg[:, ax], power=power)
            if amp is None:
                r = None
                break
            r[f"hr_{label}"], r[f"ihr_{label}"] = harmonic_ratio(amp, odd)
        if r:
            per.append(r)
    per = pd.DataFrame(per)
    # §16.2 regularity, and the trunk acceleration's RMS
    out = []
    stride_rows = np.median([b - a for a, b in strides])
    steady = acc[int(strides[0][0]):int(strides[-1][1])]
    for ax, label in ((0, "ML"), (1, "AP"), (2, "VT")):
        reg = regularity(unbiased_autocorr(steady[:, ax],
                                           int(1.6 * stride_rows)), stride_rows)
        for k in ("step_regularity", "stride_regularity", "symmetry"):
            out.append(row(f"{k}_{label}", reg[k], note=f"from {sig['name']}"))
        out.append(row(f"trunk_acc_rms_{label}",
                       np.sqrt(np.nanmean(steady[:, ax] ** 2)),
                       note=f"from {sig['name']}"))
    # §15 multiscale entropy of the trunk acceleration, with a white-noise
    # reference through the identical calculation
    seg = steady[:int(MSE_MAX_S * FS)]
    curves = {label: rcmse(seg[:, ax]) for ax, label in
              ((0, "ML"), (1, "AP"), (2, "VT"))}
    curves["white_noise"] = rcmse(rng_for("mse").normal(size=len(seg)))
    mse = pd.DataFrame(curves, index=pd.RangeIndex(1, MSE_SCALES + 1,
                                                   name="scale"))
    for label in ("ML", "AP", "VT"):
        out.append(row(f"mse_area_{label}", np.nansum(mse[label])))
    # §16.1 HR: median (unbounded); iHR: mean with a block-bootstrap CI
    for col in [c for c in per if c.startswith(("hr_", "ihr_"))]:
        res = series_summary(per[col].to_numpy(float), ("mean",),
                             tag=f"{t.stem}{col}")["mean"]
        out.append(row(f"harmonic_ratio_{col}", res["value"], limb, "mean",
                       (res["ci_low"], res["ci_high"]), res["n"],
                       res["block"],
                       note=f"per stride from {sig['name']}, "
                            f"k^{sig['power']} weighting"))
        out.append(row(f"harmonic_ratio_{col}", per[col].median(), limb,
                       "median", n=len(per)))
    return out, mse


# %% ==========================================================================
# §17  CONTROL OF STEPPING: GEM AND FOOT PLACEMENT  (Word document §17, Fig. 23)
# =============================================================================
# §17.1 GOAL-EQUIVALENT MANIFOLD (Dingwell, John & Cusumano 2010). On a belt
# at fixed speed the task goal is L/T = belt speed. Stride time T and length
# L are divided by their means (both dimensionless, goal line = diagonal);
# each stride's deviation splits into
#     along the line  (T^ + L^)/sqrt2  goal-EQUIVALENT: speed unchanged
#     across it       (L^ - T^)/sqrt2  goal-RELEVANT:   speed error
# Reported: SD (% of the mean stride), lag-1 autocorrelation, DFA alpha of
# each. Healthy walkers correct the goal-relevant part hard (alpha < 0.5)
# and leave the goal-equivalent part alone (alpha ~ 1).
# §17.2 FOOT PLACEMENT (Wang & Srinivasan 2014): regress where the other
# foot lands (across the belt) on the system CoM's position and velocity at
# this foot's mid-stance, all relative to this (stance) ankle:
#     y_foot = b0 + b1 y_CoM + b2 v_CoM ;  R2 = how much of foot placement
# the body's state explains (healthy ML > 0.8).

def lag1(x):
    x = np.asarray(x, float)
    x = x[np.isfinite(x)] - np.nanmean(x)
    s = np.sum(x ** 2)
    return float(np.sum(x[:-1] * x[1:]) / s) if s > 0 else np.nan


def gem_decompose(stride_time, stride_length):
    T = np.asarray(stride_time, float)
    L = np.asarray(stride_length, float)
    ok = np.isfinite(T) & np.isfinite(L)
    Th, Lh = T[ok] / T[ok].mean() - 1, L[ok] / L[ok].mean() - 1
    par = (Th + Lh) / np.sqrt(2)
    perp = (Lh - Th) / np.sqrt(2)
    return dict(v_star=float(L[ok].mean() / T[ok].mean()), n=int(ok.sum()),
                parallel=par, perpendicular=perp,
                parallel_sd=float(np.std(par, ddof=1)),
                perpendicular_sd=float(np.std(perp, ddof=1)),
                parallel_lag1=lag1(par), perpendicular_lag1=lag1(perp))


def foot_placement_model(com_position, com_velocity, foot):
    z, v, y = (np.asarray(a, float) for a in (com_position, com_velocity, foot))
    ok = np.isfinite(z) & np.isfinite(v) & np.isfinite(y)
    if ok.sum() < 10:
        return dict(r2=np.nan, intercept=np.nan, gain_position=np.nan,
                    gain_velocity=np.nan, n=int(ok.sum()))
    X = np.column_stack([np.ones(ok.sum()), z[ok], v[ok]])
    beta = np.linalg.lstsq(X, y[ok], rcond=None)[0]
    ss = np.sum((y[ok] - y[ok].mean()) ** 2)
    r2 = 1 - np.sum((y[ok] - X @ beta) ** 2) / ss if ss > 0 else np.nan
    return dict(r2=float(r2), intercept=float(beta[0]),
                gain_position=float(beta[1]), gain_velocity=float(beta[2]),
                n=int(ok.sum()))


def foot_placement_rows(t, limb, s):
    s = s[np.isfinite(s["to"]) & np.isfinite(s["contra_hs"])]
    if len(s) < 20:
        return []
    lat = np.r_[t.lateral, 0]
    mid = (s["hs"] + 0.5 * (s["to"] - s["hs"])).to_numpy()   # mid-stance
    ank = at(t.ankle[limb], mid)
    z = (at(t.com, mid) - ank) @ lat
    v = at(t.com_vel, mid) @ lat
    y = (at(t.heel[OTHER[limb]], s["contra_hs"].to_numpy()) - ank) @ lat
    fit = foot_placement_model(z, v, y)
    ok = np.isfinite(z) & np.isfinite(v) & np.isfinite(y)
    z, v, y = z[ok], v[ok], y[ok]
    block = optimal_block_length(y)
    ci = bootstrap_rows(len(y), lambda i: foot_placement_model(
        z[i], v[i], y[i])["r2"], block, n_boot=400, tag=f"{t.stem}fp{limb}")
    return [row("foot_placement_r2", fit["r2"], limb, ci=ci, n=fit["n"],
                block=block,
                note=f"gains {fit['gain_position']:+.2f} (position), "
                     f"{fit['gain_velocity']:+.3f} s (velocity)")]


def series_for(t, limb, n):
    """One foot's consecutive steady strides, cut to n (§12.4)."""
    s = steady_steps(t, limb)
    return s.iloc[:n] if n else s


def variability_rows(t, n):
    """§14 DFA, §15 sample entropy, §17 GEM and foot placement, per foot."""
    out = []
    for limb in LIMBS:
        s = series_for(t, limb, n)
        for r in stride_series_metrics(s[[c for c in SERIES if c in s]],
                                       tag=f"{t.stem}{limb}"):
            if r["metric"] == "dfa_alpha":
                note = (f"boxes {r['boxes']}, R2 {r['r2']:.3f}, GPH "
                        f"{r['gph']:.2f}, shuffled {r['shuffled']:.2f}, "
                        f"bias {r['bias']:+.3f}"
                        + (", metronome-paced" if r["cued"] else ""))
                out.append(row(f"dfa_alpha_{r['series']}", r["value"], limb,
                               ci=(r["ci_low"], r["ci_high"]), n=r["n"],
                               note=note))
            else:
                sweep = ", ".join(f"{k} {r[k]:.3f}" for k in r
                                  if k.startswith("r0"))
                out.append(row(f"sample_entropy_{r['series']}", r["value"],
                               limb, n=r["n"],
                               note=f"m {ENT_M}, r {ENT_R}; sweep " + sweep))
        g = gem_decompose(s["stride_s"].to_numpy(float),
                          s["stride_length_m"].to_numpy(float))
        if g["n"] >= 4 * DFA_MIN_BOX:
            for comp in ("parallel", "perpendicular"):
                r = series_summary(g[comp], ("sd",),
                                   tag=f"{t.stem}gem{comp}{limb}")["sd"]
                out.append(row(f"gem_{comp}_sd_pct", 100 * r["value"], limb,
                               ci=(100 * r["ci_low"], 100 * r["ci_high"]),
                               n=r["n"], block=r["block"],
                               note="% of the mean stride"))
                out.append(row(f"gem_{comp}_lag1", g[f"{comp}_lag1"], limb,
                               n=g["n"]))
                a = dfa_alpha(g[comp])
                lo, hi, _, _ = dfa_interval(a, g["n"],
                                            tag=f"{t.stem}{comp}{limb}")
                out.append(row(f"gem_{comp}_dfa_alpha", a, limb, ci=(lo, hi),
                               n=g["n"]))
        out.extend(foot_placement_rows(t, limb, s))
    return out


# %% ==========================================================================
# §18  LOCAL DYNAMIC STABILITY                       (Word document §18, Fig. 24)
# =============================================================================
# How fast do two nearly identical states of the walker drift apart?
# Rosenstein et al. (1993), as in Bruijn's LocalDynamicStability toolbox:
#   1. 150 consecutive strides, the whole block time-normalised to 100
#      samples per stride (cubic spline)
#   2. delay embedding: X(i) = [x(i), x(i+10), ..., x(i+40)]  (tau 10, dE 5)
#   3. every point's nearest neighbour more than half a stride away in time,
#      both followed for 10 strides; mean ln(distance) = divergence curve
#   4. lambda_S = slope over 0-0.5 stride, lambda_L = slope over 4-10
# Larger lambda = faster divergence = less locally stable.
# lambda depends on the number of strides, so EVERY window is 150 strides; a
# long trial gives up to 4 windows and the trial value is their mean.
# INTERVAL: the curve is kept per reference stride, blocks of strides are
# resampled and the curve refitted (how much lambda depends on which strides
# were walked; it does not include the method's own bias).

def lds_state_space(signals, hs, dE, scale):
    n = len(hs) - 1
    a, b = int(round(hs[0])), int(round(hs[-1]))
    block = signals[a:b + 1]
    if not np.isfinite(block).all():
        return None, None
    n_samples = n * LDS_SAMPLES_PER_STRIDE
    t_new = np.linspace(a, b, n_samples)
    norm = CubicSpline(np.arange(a, b + 1), block, axis=0)(t_new)       # 1
    if scale == "zscore":
        norm = norm / norm.std(axis=0)
    d = dE or 1
    rows = n_samples - LDS_TAU * (d - 1)
    X = np.hstack([norm[k * LDS_TAU:k * LDS_TAU + rows] for k in range(d)])  # 2
    stride = np.clip(np.searchsorted(hs, t_new[:rows], side="right") - 1,
                     0, n - 1)
    return X, stride


def divergence(X, theiler, n_lags, group=None, n_groups=1):
    """Rosenstein divergence (step 3), ln(distance) summed per group of
    reference points (a group = a stride), so strides can be resampled."""
    n = len(X)
    group = np.zeros(n, int) if group is None else group
    _, nn = cKDTree(X).query(X, k=2 * theiler + 2)
    far = np.abs(nn - np.arange(n)[:, None]) > theiler
    nbr = nn[np.arange(n), np.argmax(far, axis=1)]
    sums = np.zeros((n_groups, n_lags))
    counts = np.zeros((n_groups, n_lags))
    lags = np.arange(n_lags)
    for c in range(0, n, 400):
        i = np.arange(c, min(c + 400, n))
        a, b = i[:, None] + lags, nbr[i][:, None] + lags
        ok = (a < n) & (b < n)
        d = np.linalg.norm(X[np.minimum(a, n - 1)] - X[np.minimum(b, n - 1)],
                           axis=2)
        ok &= d > 0
        np.add.at(sums, group[i], np.where(ok, np.log(np.where(ok, d, 1.0)),
                                           0.0))
        np.add.at(counts, group[i], ok)
    return sums, counts


def fit_lambdas(curve):
    """Step 4: slopes over the two fit windows (indexed as lds_calc.m)."""
    out = {}
    for tag, (w0, w1) in LDS_FIT.items():
        i0 = max(int(round(w0 * LDS_SAMPLES_PER_STRIDE)), 1)
        i1 = min(int(round(w1 * LDS_SAMPLES_PER_STRIDE)), len(curve))
        tt = np.arange(i0, i1 + 1) / LDS_SAMPLES_PER_STRIDE
        y = curve[i0 - 1:i1]
        ok = np.isfinite(y)
        if ok.sum() < 3:
            out[tag], out[f"R2_{tag}"] = np.nan, np.nan
            continue
        out[tag] = float(np.polyfit(tt[ok], y[ok], 1)[0])
        out[f"R2_{tag}"] = float(np.corrcoef(tt[ok], y[ok])[0, 1] ** 2)
    return out


def window_lambda(signals, hs, dE, scale, tag=""):
    X, stride = lds_state_space(signals, hs, dE, scale)
    if X is None:
        return None
    n = len(hs) - 1
    sums, counts = divergence(X, LDS_SAMPLES_PER_STRIDE // 2,
                              LDS_WS * LDS_SAMPLES_PER_STRIDE, stride, n)
    with np.errstate(invalid="ignore", divide="ignore"):
        curve = sums.sum(0) / counts.sum(0)
        per_stride = sums / counts
    fit = fit_lambdas(curve)
    # block length from how correlated the strides' own short-term slopes are
    i1 = int(LDS_FIT["S"][1] * LDS_SAMPLES_PER_STRIDE)
    tt = np.arange(1, i1 + 1) / LDS_SAMPLES_PER_STRIDE
    slopes = [np.polyfit(tt, c[:i1], 1)[0] if np.isfinite(c[:i1]).all()
              else np.nan for c in per_stride]
    block = optimal_block_length(np.array(slopes))
    reps = {"S": [], "L": []}
    for k in block_indices(n, block, LDS_N_BOOT, rng_for(tag)):
        with np.errstate(invalid="ignore", divide="ignore"):
            f = fit_lambdas(sums[k].sum(0) / counts[k].sum(0))
        reps["S"].append(f["S"])
        reps["L"].append(f["L"])
    return dict(curve=curve, lambda_S=fit["S"], lambda_L=fit["L"],
                R2_S=fit["R2_S"], R2_L=fit["R2_L"], block=block,
                reps_S=np.array(reps["S"]), reps_L=np.array(reps["L"]),
                first_row=hs[0], last_row=hs[-1])


def lds_rows(t, limb):
    hs_rows = series_for(t, limb, None)["hs"].to_numpy()
    n_win = min(LDS_MAX_WINDOWS, (len(hs_rows) - 1) // LDS_N_STRIDES)
    out, windows, curves = [], [], {}
    if n_win < 1:
        print(f"  LDS: only {len(hs_rows) - 1} strides, fewer than the "
              f"{LDS_N_STRIDES} one window needs -- skipped")
        return out, pd.DataFrame(), curves
    for name, (sigs, dE, scale) in LDS_STATE_SPACES.items():
        cols = [t.kin_walker(col) for col, _ in sigs]
        if any(c is None for c in cols):
            continue
        signals = np.column_stack([c[:, ax] for c, (_, ax) in zip(cols, sigs)])
        results = []
        for w in range(n_win):
            hs = hs_rows[w * LDS_N_STRIDES:(w + 1) * LDS_N_STRIDES + 1]
            r = window_lambda(signals, hs, dE, scale, tag=f"{t.stem}{name}{w}")
            if r is None:
                continue
            results.append(r)
            windows.append(dict(state_space=name, window=w,
                                n_strides=LDS_N_STRIDES,
                                first_row=r["first_row"],
                                last_row=r["last_row"],
                                lambda_S=r["lambda_S"], lambda_L=r["lambda_L"],
                                R2_S=r["R2_S"], R2_L=r["R2_L"],
                                block=r["block"]))
        if not results:
            continue
        curves[name] = np.nanmean([r["curve"] for r in results], axis=0)
        for tag in ("S", "L"):
            value = float(np.mean([r[f"lambda_{tag}"] for r in results]))
            lo, hi = percentile_ci(np.nanmean([r[f"reps_{tag}"]
                                               for r in results], axis=0))
            out.append(row(f"lds_lambda_{tag}_{name}", value, limb,
                           ci=(lo, hi), n=LDS_N_STRIDES,
                           block=float(np.mean([r["block"] for r in results])),
                           note=f"mean of {len(results)} windows of "
                                f"{LDS_N_STRIDES} strides"
                                + (", primary" if name == LDS_PRIMARY else "")))
    return out, pd.DataFrame(windows), curves


# %% ==========================================================================
# §19  RESULTS TABLES                              (Word document §19, Fig. 25)
# =============================================================================
# Two shapes of the same numbers:
#   LONG  (all_trials_summary.csv): one row per number with its CI, N,
#         block length and a note -- the shape a mixed model needs
#         (participant random, condition fixed).
#   WIDE  (all_trials_metrics.xlsx/.csv): one row per trial, one column per
#         metric -- the shape to read. Column names:
#           step_width_m      mean of all steady steps (both feet)
#           step_width_m_sd   SD;  _cv  CV (%)
#           step_width_m_L/_R each foot's mean, or for DFA / entropy / GEM /
#                             foot placement the value from that foot's series
# The workbook also colours every trial against the healthy reference
# ranges below (Word document Table 19.1), and has a dictionary of every
# column with the document section that defines it.

def summarise(t, series_length):
    """Every per-trial number of §13-§18 in the long form."""
    rows = trial_descriptors(t) + step_summaries(t) + symmetry_rows(t)
    rows += variability_rows(t, series_length)
    trunk, mse = trunk_rows(t, LIMB_FOR_TRUNK, series_length)
    rows += trunk
    extras = dict(mse=mse, lds_windows=pd.DataFrame(), lds_curves={})
    if RUN_LDS:
        lds, extras["lds_windows"], extras["lds_curves"] = lds_rows(
            t, LIMB_FOR_TRUNK)
        rows += lds
    out = pd.DataFrame(rows)
    out.insert(0, "trial", t.stem)
    parts = t.stem.split("_")
    out.insert(1, "participant", parts[0])
    out.insert(2, "condition", parts[1] if len(parts) > 1 else "")
    out["series_length"] = series_length or np.nan
    return out, extras


# --- what every column means: (family, description, unit, Word section) ---------
CATALOG = {
    "belt_speed_ms": ("Trial", "Belt speed, measured from the stance heels", "m/s", "5.2"),
    "system_mass_kg": ("Trial", "Body + load, weighed in the quiet standing", "kg", "5.5"),
    "load_kg": ("Trial", "Carried load: system mass minus body mass", "kg", "5.5"),
    "load_offset_y_m": ("Trial", "Load CoM forward of Trunk_Position (negative = behind)", "m", "5.6"),
    "load_offset_x_m": ("Trial", "Load CoM to the right of Trunk_Position (0 when centred)", "m", "5.6"),
    "load_offset_z_m": ("Trial", "Load CoM above Trunk_Position (assumed)", "m", "5.6"),
    "quiet_standing_cv": ("Trial", "CV of total vertical force in the quiet standing", "-", "5.5"),
    "steps_steady": ("Trial", "Steps after the warm-up", "steps", "5.1"),
    "events_trusted_grf_pct": ("Trial", "Events from trusted force plates", "%", "4"),
    "events_zeni_pct": ("Trial", "Events from Zeni (kinematic fallback)", "%", "4"),
    "events_interpolated_pct": ("Trial", "Events interpolated (last resort)", "%", "4"),
    "stances_kinetics_pct": ("Trial", "Steady stances trusted for kinetics", "%", "8"),
    "swings_without_mfc_pct": ("Trial", "Swings with no valid MFC event (non-MTC)", "%", "11"),
    "mfc_below_belt_pct": ("Trial", "Swings whose MFC is below the belt (pose error)", "%", "11"),
    "handrail_steps_pct": ("Trial", "Steady steps with handrail contact", "%", "5.4"),
    "grf_vs_marker_acc_r_X": ("Trial", "GRF vs marker CoM acceleration, Theia X", "r", "5.7"),
    "grf_vs_marker_acc_r_Y": ("Trial", "GRF vs marker CoM acceleration, Theia Y", "r", "5.7"),
    "grf_vs_marker_acc_r_Z": ("Trial", "GRF vs marker CoM acceleration, vertical", "r", "5.7"),
    "stride_s": ("Spatiotemporal", "Stride time", "s", "6"),
    "stance_s": ("Spatiotemporal", "Stance time", "s", "6"),
    "swing_s": ("Spatiotemporal", "Swing time", "s", "6"),
    "step_s": ("Spatiotemporal", "Step time", "s", "6"),
    "cadence_spm": ("Spatiotemporal", "Cadence", "steps/min", "6"),
    "stance_pct": ("Spatiotemporal", "Stance, % of stride", "%", "6"),
    "double_support_pct": ("Spatiotemporal", "Double support, % of stride", "%", "6"),
    "single_support_pct": ("Spatiotemporal", "Single support, % of stride", "%", "6"),
    "step_length_m": ("Spatiotemporal", "Step length (heel to heel, along the belt)", "m", "6"),
    "step_width_m": ("Spatiotemporal", "Step width (heel to heel, across the belt)", "m", "6"),
    "stride_length_m": ("Spatiotemporal", "Stride length over the belt", "m", "6"),
    "stride_speed_ms": ("Spatiotemporal", "Stride speed over the belt", "m/s", "6"),
    "walk_ratio": ("Spatiotemporal", "Walk ratio: step length / cadence", "m/(steps/min)", "13.3"),
    "trunk_lean_mean_deg": ("Posture", "Trunk (thorax) lean, mean over the stride", "deg", "7"),
    "trunk_lean_rom_deg": ("Posture", "Trunk lean, range within the stride", "deg", "7"),
    "pelvis_tilt_mean_deg": ("Posture", "Pelvis tilt, mean over the stride", "deg", "7"),
    "hip_rom_deg": ("Posture", "Hip sagittal range of motion", "deg", "7"),
    "knee_rom_deg": ("Posture", "Knee sagittal range of motion", "deg", "7"),
    "ankle_rom_deg": ("Posture", "Ankle sagittal range of motion", "deg", "7"),
    "f1_bw": ("Kinetics", "Vertical GRF first peak (weight acceptance)", "BW", "8"),
    "trough_bw": ("Kinetics", "Vertical GRF mid-stance trough", "BW", "8"),
    "f2_bw": ("Kinetics", "Vertical GRF second peak (push-off)", "BW", "8"),
    "f1_tw": ("Kinetics", "Vertical GRF first peak, per total weight", "TW", "8"),
    "f2_tw": ("Kinetics", "Vertical GRF second peak, per total weight", "TW", "8"),
    "loading_rate_bw_per_s": ("Kinetics", "Loading rate, 20-80% of F1", "BW/s", "8"),
    "peak_braking_bw": ("Kinetics", "Peak braking force", "BW", "8"),
    "peak_propulsive_bw": ("Kinetics", "Peak propulsive force", "BW", "8"),
    "braking_impulse_bw_s": ("Kinetics", "Braking impulse (negative)", "BW s", "8"),
    "propulsive_impulse_bw_s": ("Kinetics", "Propulsive impulse", "BW s", "8"),
    "vertical_impulse_bw_s": ("Kinetics", "Vertical impulse", "BW s", "8"),
    "braking_impulse_tw_s": ("Kinetics", "Braking impulse, per total weight", "TW s", "8"),
    "propulsive_impulse_tw_s": ("Kinetics", "Propulsive impulse, per total weight", "TW s", "8"),
    "cop_ap_range_mm": ("Kinetics", "CoP excursion along the belt", "mm", "8"),
    "cop_ml_range_mm": ("Kinetics", "CoP excursion across the belt", "mm", "8"),
    "free_moment_peak_nm": ("Kinetics", "Peak |free moment|", "N m", "8"),
    "collision_j_kg": ("CoM work", "Collision: negative work, heel strike to other toe-off", "J/kg", "9"),
    "rebound_j_kg": ("CoM work", "Rebound: positive work in single support", "J/kg", "9"),
    "preload_j_kg": ("CoM work", "Preload: negative work in single support", "J/kg", "9"),
    "push_off_j_kg": ("CoM work", "Push-off: positive work, other heel strike to toe-off", "J/kg", "9"),
    "com_work_pos_j_kg": ("CoM work", "Positive work over the stance", "J/kg", "9"),
    "com_work_neg_j_kg": ("CoM work", "Negative work over the stance", "J/kg", "9"),
    "mos_ml_contact": ("Stability", "ML margin of stability at heel strike", "m", "10"),
    "mos_ml_min": ("Stability", "ML margin of stability, single-support minimum", "m", "10"),
    "mos_ap_contact": ("Stability", "AP margin of stability at heel strike", "m", "10"),
    "mos_ap_min": ("Stability", "AP margin of stability, single-support minimum", "m", "10"),
    "mfc_m": ("Stability", "Minimum foot clearance above the belt", "m", "11"),
    "moi_peak_mm": ("Stability", "Margin of instability, peak in swing", "mm", "11"),
    "tri_s": ("Stability", "Trip risk integral", "s", "11"),
    "harmonic_ratio_hr_AP": ("Trunk", "Harmonic ratio AP (median over strides)", "-", "16.1"),
    "harmonic_ratio_hr_VT": ("Trunk", "Harmonic ratio vertical (median)", "-", "16.1"),
    "harmonic_ratio_hr_ML": ("Trunk", "Harmonic ratio ML (median)", "-", "16.1"),
    "harmonic_ratio_ihr_AP": ("Trunk", "Improved harmonic ratio AP (mean)", "%", "16.1"),
    "harmonic_ratio_ihr_VT": ("Trunk", "Improved harmonic ratio vertical (mean)", "%", "16.1"),
    "harmonic_ratio_ihr_ML": ("Trunk", "Improved harmonic ratio ML (mean)", "%", "16.1"),
    "step_regularity_AP": ("Trunk", "Step regularity AP", "-", "16.2"),
    "step_regularity_VT": ("Trunk", "Step regularity vertical", "-", "16.2"),
    "step_regularity_ML": ("Trunk", "Step regularity ML", "-", "16.2"),
    "stride_regularity_AP": ("Trunk", "Stride regularity AP", "-", "16.2"),
    "stride_regularity_VT": ("Trunk", "Stride regularity vertical", "-", "16.2"),
    "stride_regularity_ML": ("Trunk", "Stride regularity ML", "-", "16.2"),
    "symmetry_AP": ("Trunk", "Step / stride regularity, AP", "-", "16.2"),
    "symmetry_VT": ("Trunk", "Step / stride regularity, vertical", "-", "16.2"),
    "symmetry_ML": ("Trunk", "Step / stride regularity, ML", "-", "16.2"),
    "trunk_acc_rms_AP": ("Trunk", "Trunk acceleration RMS, AP", "m/s2", "16"),
    "trunk_acc_rms_VT": ("Trunk", "Trunk acceleration RMS, vertical", "m/s2", "16"),
    "trunk_acc_rms_ML": ("Trunk", "Trunk acceleration RMS, ML", "m/s2", "16"),
    "mse_area_AP": ("Trunk", "Multiscale entropy area, scales 1-30, AP", "-", "15.2"),
    "mse_area_VT": ("Trunk", "Multiscale entropy area, vertical", "-", "15.2"),
    "mse_area_ML": ("Trunk", "Multiscale entropy area, ML", "-", "15.2"),
    "foot_placement_r2": ("Control", "ML foot placement explained by CoM state (R2)", "-", "17.2"),
    "gem_perpendicular_sd_pct": ("Control", "GEM goal-relevant SD", "% stride", "17.1"),
    "gem_parallel_sd_pct": ("Control", "GEM goal-equivalent SD", "% stride", "17.1"),
    "gem_perpendicular_lag1": ("Control", "GEM goal-relevant lag-1 autocorrelation", "-", "17.1"),
    "gem_parallel_lag1": ("Control", "GEM goal-equivalent lag-1 autocorrelation", "-", "17.1"),
    "gem_perpendicular_dfa_alpha": ("Control", "GEM goal-relevant DFA alpha", "-", "17.1"),
    "gem_parallel_dfa_alpha": ("Control", "GEM goal-equivalent DFA alpha", "-", "17.1"),
}
FAMILIES = ["Trial", "Spatiotemporal", "Stability", "Kinetics", "CoM work",
            "Posture", "Symmetry", "DFA", "Entropy", "Trunk", "Control",
            "LDS"]
KEY_METRICS = [
    "load_kg", "system_mass_kg", "belt_speed_ms", "steps_steady",
    "events_trusted_grf_pct", "stride_s", "cadence_spm", "stance_pct",
    "double_support_pct", "step_length_m", "step_width_m", "step_width_m_cv",
    "stride_length_m_cv", "mos_ml_contact", "mos_ml_min", "mos_ap_min",
    "mfc_m", "mfc_m_sd", "tri_s", "f1_bw", "f2_bw", "loading_rate_bw_per_s",
    "braking_impulse_bw_s", "propulsive_impulse_bw_s", "collision_j_kg",
    "push_off_j_kg", "trunk_lean_mean_deg", "knee_rom_deg",
    "dfa_alpha_stride_s_L", "dfa_alpha_stride_s_R",
    "dfa_alpha_step_width_m_L", "dfa_alpha_step_width_m_R",
    "foot_placement_r2_L", "foot_placement_r2_R", "harmonic_ratio_ihr_AP",
    "lds_lambda_S_trunkVel_AP"]
ID = ["participant", "condition", "trial"]

# --- HEALTHY REFERENCE RANGES (Word document §19.2, Table 19.1) ------------------
# column: (low, high, population / conditions, source, confidence)
# confidence "confirmed" = checked against the source's abstract or a review;
# "verify" = from the gait literature at large, check before publication.
# These are healthy adults WITHOUT load at about 1.2-1.4 m/s unless stated.
# They are sanity ranges, not norms: most depend on speed, footwear, method
# and, for the variability measures, on the number of strides.
REFERENCE_RANGES = {
    "stride_s": (1.00, 1.15, "healthy adults ~1.3 m/s (fixed at 1.11 s by the 108-bpm metronome here)", "Perry & Burnfield 2010; Oberg et al. 1993", "verify"),
    "cadence_spm": (104, 120, "healthy adults ~1.3 m/s", "Oberg et al. 1993", "verify"),
    "stance_pct": (58, 62, "healthy adults, % of gait cycle", "Perry & Burnfield 2010", "verify"),
    "double_support_pct": (18, 24, "healthy adults, both periods, % of gait cycle", "Perry & Burnfield 2010", "verify"),
    "single_support_pct": (38, 42, "healthy adults, % of gait cycle", "Perry & Burnfield 2010", "verify"),
    "step_length_m": (0.65, 0.80, "healthy adults ~1.3 m/s (scales with stature)", "Oberg et al. 1993", "verify"),
    "step_width_m": (0.08, 0.15, "healthy young adults, treadmill, heel to heel", "Owings & Grabiner 2004", "verify"),
    "stride_s_cv": (1.0, 3.0, "healthy young adults", "Hausdorff 2005", "verify"),
    "stride_length_m_cv": (1.0, 3.0, "healthy young adults", "Hausdorff 2005", "verify"),
    "step_width_m_sd": (0.010, 0.030, "healthy young adults, treadmill", "Owings & Grabiner 2004", "verify"),
    "walk_ratio": (0.0055, 0.0072, "healthy adults, preferred speed (0.0063 +- 0.0007)", "Sekiya & Nagasaki 1998", "confirmed"),
    "hip_rom_deg": (38, 50, "healthy adults, sagittal", "Winter 2009; Kadaba et al. 1990", "verify"),
    "knee_rom_deg": (55, 68, "healthy adults, sagittal", "Winter 2009; Kadaba et al. 1990", "verify"),
    "ankle_rom_deg": (22, 35, "healthy adults, sagittal", "Winter 2009; Kadaba et al. 1990", "verify"),
    "f1_bw": (1.05, 1.30, "healthy adults, treadmill (e.g. 1.19 BW)", "Winter 2009; treadmill GRF studies", "confirmed"),
    "trough_bw": (0.60, 0.85, "healthy adults ~1.3 m/s", "Winter 2009", "verify"),
    "f2_bw": (1.00, 1.25, "healthy adults, treadmill (e.g. 1.14 BW)", "Winter 2009; treadmill GRF studies", "confirmed"),
    "f1_tw": (1.05, 1.30, "unloaded pattern (per total weight)", "Birrell et al. 2007", "verify"),
    "f2_tw": (1.00, 1.25, "unloaded pattern (per total weight)", "Birrell et al. 2007", "verify"),
    "loading_rate_bw_per_s": (5.0, 15.0, "shod walking, mean slope 20-80% of F1", "gait literature; method dependent", "verify"),
    "peak_braking_bw": (0.15, 0.25, "healthy adults ~1.3 m/s", "Winter 2009", "verify"),
    "peak_propulsive_bw": (0.15, 0.25, "healthy adults ~1.3 m/s", "Winter 2009", "verify"),
    "propulsive_impulse_bw_s": (0.02, 0.04, "healthy adults", "Revi et al. 2020", "verify"),
    "braking_impulse_bw_s": (-0.04, -0.02, "healthy adults", "Revi et al. 2020", "verify"),
    "vertical_impulse_bw_s": (0.50, 0.60, "physics: each limb carries ~half the weight over a stride of 1.0-1.2 s", "Newton", "confirmed"),
    "cop_ap_range_mm": (150, 220, "healthy adults (~60-75% of foot length)", "gait literature", "verify"),
    "cop_ml_range_mm": (15, 40, "healthy adults", "gait literature", "verify"),
    "free_moment_peak_nm": (2, 8, "healthy adults walking", "Holden & Cavanagh 1991; Li et al. 2001", "verify"),
    "push_off_j_kg": (0.10, 0.25, "healthy adults 1.25 m/s, per step", "Donelan et al. 2002", "verify"),
    "collision_j_kg": (-0.25, -0.08, "healthy adults 1.25 m/s, per step", "Donelan et al. 2002", "verify"),
    "mos_ml_contact": (0.03, 0.10, "healthy adults, ankle/malleolus boundary (boot edge adds ~0.03-0.05 m); treadmill > overground", "Hof et al. 2005; Rosenblum et al. 2021", "verify"),
    "mos_ml_min": (0.02, 0.08, "healthy adults, ankle boundary", "Hof et al. 2005", "verify"),
    "mfc_m": (0.010, 0.030, "healthy young adults (1-2 cm overground; treadmill studies up to 3 cm)", "Begg et al. 2007; Barrett et al. 2010", "confirmed"),
    "mfc_m_sd": (0.002, 0.006, "healthy young adults", "Begg et al. 2007; Barrett et al. 2010", "verify"),
    "harmonic_ratio_hr_AP": (3.0, 4.0, "healthy young adults, trunk", "Menz et al. 2003; Lowry et al. 2012", "confirmed"),
    "harmonic_ratio_hr_VT": (3.0, 4.0, "healthy young adults, trunk", "Menz et al. 2003; Lowry et al. 2012", "confirmed"),
    "harmonic_ratio_hr_ML": (2.0, 2.7, "healthy young adults, trunk", "Menz et al. 2003; Lowry et al. 2012", "confirmed"),
    "harmonic_ratio_ihr_AP": (85, 97, "healthy young adults, sacrum", "Pasciuto et al. 2017", "verify"),
    "harmonic_ratio_ihr_VT": (85, 97, "healthy young adults, sacrum", "Pasciuto et al. 2017", "verify"),
    "harmonic_ratio_ihr_ML": (70, 90, "healthy young adults, sacrum", "Pasciuto et al. 2017", "verify"),
    "step_regularity_VT": (0.70, 0.95, "healthy adults, trunk", "Moe-Nilssen & Helbostad 2004; Kobsar et al. 2014", "verify"),
    "stride_regularity_VT": (0.75, 0.95, "healthy adults, trunk", "Moe-Nilssen & Helbostad 2004; Kobsar et al. 2014", "verify"),
    "step_regularity_AP": (0.60, 0.90, "healthy adults, trunk", "Moe-Nilssen & Helbostad 2004; Kobsar et al. 2014", "verify"),
    "stride_regularity_AP": (0.65, 0.90, "healthy adults, trunk", "Moe-Nilssen & Helbostad 2004; Kobsar et al. 2014", "verify"),
    "dfa_alpha_stride_s": (0.75, 0.90, "healthy adults, uncued; METRONOME-PACED walking gives < 0.5", "Hausdorff et al. 1996; Terrier et al. 2005", "confirmed"),
    "dfa_alpha_stride_length_m": (0.60, 0.90, "healthy young adults, treadmill (persistent)", "Dingwell et al. 2010", "confirmed"),
    "dfa_alpha_stride_speed_ms": (0.20, 0.40, "healthy young adults, treadmill (anti-persistent)", "Dingwell et al. 2010; Terrier 2012", "confirmed"),
    "dfa_alpha_step_width_m": (0.60, 0.90, "healthy young adults, treadmill (persistent)", "Dingwell & Cusumano 2015", "verify"),
    "gem_perpendicular_dfa_alpha": (0.20, 0.50, "healthy young adults, treadmill", "Dingwell et al. 2010", "verify"),
    "gem_parallel_dfa_alpha": (0.70, 1.00, "healthy young adults, treadmill", "Dingwell et al. 2010", "verify"),
    "foot_placement_r2": (0.60, 0.90, "healthy adults, ML, mid-stance (> 0.8 with pelvis state)", "Wang & Srinivasan 2014", "confirmed"),
    "lds_lambda_S_trunkVel_AP": (0.40, 0.60, "healthy adults, trunk, per stride (0.50 +- 0.06); strongly method dependent", "van Schooten et al. 2011", "confirmed"),
    "symmetry_angle_stance_s": (-3, 3, "healthy adults", "Zifchock et al. 2008", "verify"),
    "symmetry_angle_step_length_m": (-3, 3, "healthy adults", "Zifchock et al. 2008", "verify"),
    "symmetry_angle_f1_bw": (-3, 3, "healthy adults", "Zifchock et al. 2008", "verify"),
}


def describe(metric):
    """(family, description, unit, section) for any metric of the long
    table, including those named after the series they are computed on."""
    if metric in CATALOG:
        return CATALOG[metric]
    for prefix, family, what, section in (
            ("dfa_alpha_", "DFA", "DFA alpha of", "14"),
            ("sample_entropy_", "Entropy",
             "Sample entropy (m 2, r 0.2 SD) of", "15.1"),
            ("symmetry_angle_", "Symmetry",
             "Symmetry angle (0 = symmetric, + = right larger) of", "13.2")):
        if metric.startswith(prefix):
            base = metric[len(prefix):]
            label = CATALOG.get(base, ("", base))[1].lower()
            return (family, f"{what} {label}",
                    "%" if family == "Symmetry" else "-", section)
    if metric.startswith("lds_lambda_"):
        tag, space = metric[len("lds_lambda_"):].split("_", 1)
        span = ("short-term (0-0.5 stride)" if tag == "S"
                else "long-term (4-10 strides)")
        return ("LDS", f"Lyapunov exponent, {span}, state space {space}",
                "1/stride", "18")
    return ("Other", metric, "", "")


def wide_column(metric, limb, statistic):
    """The wide-table column a long-table row goes to (None = not shown)."""
    if metric.startswith("harmonic_ratio_hr_"):       # unbounded -> median
        return metric if statistic == "median" else None
    if metric.startswith("harmonic_ratio_ihr_"):      # bounded -> mean
        return metric if statistic == "mean" else None
    if metric.startswith("lds_"):
        return metric
    if statistic in ("sd", "cv"):
        return f"{metric}_{statistic}" if limb == "both" else None
    if statistic in ("mean", "value", "pct"):
        return metric if limb == "both" else f"{metric}_{limb}"
    return None


def reference_for(column):
    """The healthy range for a wide column (per-foot columns share the
    range of their metric)."""
    if column in REFERENCE_RANGES:
        return REFERENCE_RANGES[column]
    for suffix in ("_L", "_R"):
        if column.endswith(suffix) and column[:-2] in REFERENCE_RANGES:
            return REFERENCE_RANGES[column[:-2]]
    return None


def wide_tables(long):
    """Long rows -> (values, CI limits, dictionary), one row per trial."""
    long = long.copy()
    long["column"] = [wide_column(m, b, s) for m, b, s in
                      zip(long["metric"], long["limb"], long["statistic"])]
    long = long.dropna(subset=["column"])
    meta = {}
    for m, b, s, c in zip(long["metric"], long["limb"], long["statistic"],
                          long["column"]):
        fam, desc, unit, sec = describe(m)
        foot = {"L": "left foot", "R": "right foot", "both": "both feet"}[b]
        if s in ("sd", "cv"):
            desc = f"{desc}, {'SD' if s == 'sd' else 'CV'} over steps ({foot})"
            unit = "%" if s == "cv" else unit
        elif s == "mean" and fam != "Trunk":
            desc = f"{desc}, mean over steps ({foot})"
        elif b in LIMBS and c.endswith(f"_{b}"):
            desc = f"{desc}, {foot}"
        ref = reference_for(c)
        meta[c] = dict(column=c, family=fam, description=desc, unit=unit,
                       limb=b, statistic=s, methods_section=sec,
                       healthy_low=ref[0] if ref else np.nan,
                       healthy_high=ref[1] if ref else np.nan,
                       healthy_source=ref[3] if ref else "", metric=m)
    # order: family, then each metric where it first appears, then mean,
    # SD, CV, left, right
    order = {f: i for i, f in enumerate(FAMILIES + ["Other"])}
    first = list(dict.fromkeys(long["metric"]))
    variant = {"both": 0, "L": 3, "R": 4}
    cols = sorted(meta, key=lambda c: (
        order[meta[c]["family"]], first.index(meta[c]["metric"]),
        {"sd": 1, "cv": 2}.get(meta[c]["statistic"],
                               variant[meta[c]["limb"]])))
    values = (long.pivot_table(index=ID, columns="column", values="value",
                               aggfunc="first").reindex(columns=cols)
              .reset_index())
    lo = long.pivot_table(index=ID, columns="column", values="ci_low",
                          aggfunc="first")
    hi = long.pivot_table(index=ID, columns="column", values="ci_high",
                          aggfunc="first")
    ci_cols = [c for c in cols if c in lo and lo[c].notna().any()]
    ci = pd.concat([lo[ci_cols].add_suffix("_lo"),
                    hi[ci_cols].add_suffix("_hi")], axis=1)
    ci = ci[[f"{c}_{e}" for c in ci_cols for e in ("lo", "hi")]].reset_index()
    values.columns.name = ci.columns.name = None
    dictionary = pd.DataFrame([meta[c] for c in cols]).drop(columns="metric")
    return values, ci, dictionary


def build_tables(folder=None):
    """Rebuild every results table from every {trial}_summary.csv in the
    output folder (so earlier runs are included)."""
    folder = Path(folder or OUTPUT_FOLDER)
    files = [f for f in sorted(folder.glob("*_summary.csv"))
             if not f.name.startswith("all_trials")]
    if not files:
        return None
    long = pd.concat([pd.read_csv(f) for f in files], ignore_index=True)
    long.to_csv(folder / "all_trials_summary.csv", index=False)
    values, ci, dictionary = wide_tables(long)
    values.to_csv(folder / "all_trials_metrics.csv", index=False)
    ci.to_csv(folder / "all_trials_metrics_ci.csv", index=False)
    dictionary.to_csv(folder / "metric_dictionary.csv", index=False)
    xlsx = folder / "all_trials_metrics.xlsx"
    try:
        write_workbook(xlsx, values, ci, dictionary)
    except ImportError:
        print("  (pip install openpyxl for the .xlsx workbook; the CSVs are "
              "written)")
        xlsx = None
    print(f"\n{len(values)} trials x {values.shape[1] - 3} metrics -> "
          f"{folder / 'all_trials_metrics.csv'}"
          + (f"\n  and {xlsx.name}" if xlsx else ""))
    return values, ci, dictionary


def write_workbook(path, values, ci, dictionary):
    """Sheets: Key metrics, one per family, Healthy ranges (every trial
    coloured green inside / orange outside the reference range), 95% CI,
    Dictionary. Row 1 = description (unit), row 2 = column name."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    info = dictionary.set_index("column")
    wb = Workbook()
    wb.remove(wb.active)
    inside = PatternFill("solid", fgColor="D9EAD3")
    outside = PatternFill("solid", fgColor="FCE5CD")

    def sheet(name, cols, frame=values, labels=None, colour=False):
        ws = wb.create_sheet(name[:31])
        cols = ID + [c for c in cols if c in frame]
        for j, c in enumerate(cols, 1):
            if labels is not None:
                label = labels.get(c, c)
            elif c in info.index:
                label = f"{info.at[c, 'description']} ({info.at[c, 'unit']})"
                if colour and reference_for(c):
                    lo_, hi_ = reference_for(c)[:2]
                    label += f"  healthy {lo_:g} to {hi_:g}"
            else:
                label = c.capitalize()
            top = ws.cell(row=1, column=j, value=label)
            top.font = Font(bold=True)
            top.alignment = Alignment(wrap_text=True, vertical="top")
            top.fill = PatternFill("solid", fgColor="DDE6F0")
            ws.cell(row=2, column=j, value=c).font = Font(
                italic=True, color="666666", size=9)
            ref = reference_for(c) if colour else None
            for i, v in enumerate(frame[c].tolist(), 3):
                if isinstance(v, float) and not np.isfinite(v):
                    v = None
                cell = ws.cell(row=i, column=j, value=v)
                if isinstance(v, float):
                    cell.number_format = "0.000" if abs(v) < 100 else "0.0"
                    if ref:
                        cell.fill = inside if ref[0] <= v <= ref[1] else outside
            ws.column_dimensions[get_column_letter(j)].width = \
                34 if c == "trial" else 16
        ws.row_dimensions[1].height = 75
        ws.freeze_panes = "D3"

    sheet("Key metrics", KEY_METRICS)
    for fam in FAMILIES + ["Other"]:
        cols = dictionary.loc[dictionary["family"] == fam, "column"].tolist()
        if cols:
            sheet(fam, cols)
    sheet("Healthy ranges", [c for c in values.columns[3:]
                             if reference_for(c)], colour=True)
    ci_labels = {c: (f"{info.at[c[:-3], 'description']}, 95% CI "
                     f"{'lower' if c.endswith('_lo') else 'upper'}")
                 for c in ci.columns if c[:-3] in info.index}
    sheet("95% CI", [c for c in ci.columns if c not in ID], ci, ci_labels)
    ws = wb.create_sheet("Dictionary")
    for j, h in enumerate(dictionary.columns, 1):
        ws.cell(row=1, column=j, value=h).font = Font(bold=True)
    for i, r in enumerate(dictionary.itertuples(index=False), 2):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j,
                    value=None if isinstance(v, float) and not np.isfinite(v)
                    else v)
    for j, w in enumerate((34, 14, 70, 12, 6, 10, 10, 10, 10, 30), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"
    wb.save(path)


# %% ==========================================================================
# §19  FIGURES                                         (Word document §19.3)
# =============================================================================
# {trial}_qa.png: is the trial sound? -- the quiet standing and the load,
#   where the events came from, clearance over the belt, the margins over
#   the trial, the GRF of trusted stances, the velocity fusion.
# {trial}_variability.png: DFA with its interval, alpha for every series,
#   the LDS divergence curves, the multiscale entropy curves.
# {trial}_mos_stride.mp4 (MAKE_VIDEO): birds-eye view of one typical stride.

COLOUR = {"L": "#2a5d9f", "R": "#a8442a"}
SOURCE_COLOUR = {"GRF": "#1b7a3d", "kinematic": "#d98b04",
                 "GRF_unverified": "#7a4fb5", "interpolated": "#c0392b"}


def qa_figure(t, path):
    fig, ax = plt.subplots(3, 2, figsize=(15, 12))
    s = t.steps

    # 1. quiet standing and the weight
    n = int(min(30 * FS_FORCE, t.n_force))
    tt = np.arange(n) / FS_FORCE
    total = t.grf["L"][:n, 2] + t.grf["R"][:n, 2]
    a = ax[0, 0]
    for b in LIMBS:
        a.plot(tt, t.grf[b][:n, 2], color=COLOUR[b], lw=0.6,
               label=f"{SIDE[b]} belt")
    a.plot(tt, total, color="k", lw=0.8, label="total")
    L = t.load
    if L["found"]:
        a.axvspan(L["start_s"], L["stop_s"], color="#f2d24b", alpha=0.4)
        a.axhline(L["weight_n"], color="k", ls=":", lw=1)
    a.set_title(f"Quiet standing: {t.mass:.1f} kg total, load "
                f"{L['load_kg']:+.1f} kg", fontsize=10)
    a.set_xlabel("s")
    a.set_ylabel("vertical GRF (N)")
    a.legend(fontsize=7, frameon=False)

    # 2. stance time over the trial, coloured by where the events came from
    a = ax[0, 1]
    src = np.where(s["hs_source"] == "GRF", s["to_source"], s["hs_source"])
    for name, c in SOURCE_COLOUR.items():
        m = src == name
        a.plot(s["hs"][m] / FS, s["stance_s"][m], ".", ms=3, color=c,
               label=f"{name} ({m.sum()})")
    a.axvline(s.loc[s["steady"], "hs"].min() / FS, color="k", lw=0.8,
              ls="--")
    a.set_title("Stance time, by event source (dashed: end of warm-up)",
                fontsize=10)
    a.set_xlabel("s")
    a.set_ylabel("stance (s)")
    a.legend(fontsize=7, frameon=False, markerscale=3)

    # 3. clearance over a few strides, with MFC
    a = ax[1, 0]
    mid = s[s["steady"] & (s["limb"] == "R")]
    if len(mid) > 6:
        r0 = int(mid["hs"].iloc[len(mid) // 2])
        r1 = int(mid["hs"].iloc[len(mid) // 2 + 4])
        rows = np.arange(r0, r1)
        for b in LIMBS:
            a.plot(rows / FS, 1000 * t.boots.low_z[b][rows], color=COLOUR[b],
                   lw=1, label=f"{SIDE[b]} lowest sole point")
        m = s[(s["mfc_row"] >= r0) & (s["mfc_row"] < r1)]
        a.plot(m["mfc_row"] / FS, 1000 * m["mfc_m"], "kv", ms=6, label="MFC")
    a.axhline(0, color="k", lw=0.8)
    a.set_ylim(-10, 120)
    a.set_title(f"Height above the belt (MFC median "
                f"{1000 * s['mfc_m'].median():.1f} mm, "
                f"{(s['mfc_m'] < 0).sum()} below the belt)", fontsize=10)
    a.set_xlabel("s")
    a.set_ylabel("mm")
    a.legend(fontsize=7, frameon=False)

    # 4. margins over the trial
    a = ax[1, 1]
    for b in LIMBS:
        m = s["limb"] == b
        a.plot(s["hs"][m] / FS, 1000 * s["mos_ml_contact"][m], ".", ms=2,
               color=COLOUR[b], label=f"ML at contact, {SIDE[b]}")
        a.plot(s["hs"][m] / FS, 1000 * s["mos_ap_min"][m], "x", ms=2,
               color=COLOUR[b], alpha=0.5, label=f"AP stance minimum, {SIDE[b]}")
    a.axhline(0, color="k", lw=0.8)
    a.set_title("Margins of stability (system CoM, boot soles)", fontsize=10)
    a.set_xlabel("s")
    a.set_ylabel("mm")
    a.legend(fontsize=7, frameon=False, markerscale=3, ncol=2)

    # 5. vertical GRF ensemble, trusted stances
    a = ax[2, 0]
    if t.waves["vgrf"]:
        W = np.array(t.waves["vgrf"])
        limb = np.array(t.waves["limb"])
        pct = np.linspace(0, 100, W.shape[1])
        for b in LIMBS:
            if (limb == b).any():
                w = W[limb == b]
                a.fill_between(pct, *np.percentile(w, [5, 95], axis=0),
                               color=COLOUR[b], alpha=0.2, lw=0)
                a.plot(pct, w.mean(0), color=COLOUR[b],
                       label=f"{SIDE[b]} (n={len(w)})")
    a.axhline(1, color="k", ls=":", lw=0.8)
    a.set_title("Vertical GRF, trusted stances (per total weight)",
                fontsize=10)
    a.set_xlabel("% stance")
    a.legend(fontsize=7, frameon=False)

    # 6. the velocity fusion, ML, a few seconds of steady walking
    a = ax[2, 1]
    r0 = int(s.loc[s["steady"], "hs"].min()) if s["steady"].any() else 0
    rows = np.arange(r0, min(r0 + 4 * FS, t.n))
    k = np.r_[t.lateral, 0]
    a.plot(rows / FS, t.com_vel_markers[rows] @ k, color="#bbbbbb", lw=1,
           label="markers, differentiated")
    a.plot(rows / FS, t.com_vel[rows] @ k, color="k", lw=1.2,
           label="fused with integrated GRF")
    r = t.fusion_check
    a.set_title("CoM velocity across the walking direction (GRF vs marker "
                f"acceleration r = {r['X']:.2f}/{r['Y']:.2f}/{r['Z']:.2f})",
                fontsize=10)
    a.set_xlabel("s")
    a.set_ylabel("m/s")
    a.legend(fontsize=7, frameon=False)

    for a in ax.flat:
        a.grid(alpha=0.15)
    fig.suptitle(t.stem, fontsize=12)
    fig.tight_layout()
    fig.savefig(path, dpi=110)
    plt.close(fig)


def variability_figure(t, summary, extras, series_length, path, limb="R"):
    fig, ax = plt.subplots(2, 2, figsize=(14, 9))

    # 1. DFA of the first uncued series, with its interval
    s = t.steps[t.steps["steady"] & (t.steps["limb"] == limb)].sort_values("hs")
    s = s.iloc[:series_length] if series_length else s
    a = ax[0, 0]
    for name in ("step_width_m", "mos_ml_contact", "stride_s"):
        x = s[name].to_numpy(float) if name in s else np.array([])
        if np.isfinite(x).mean() > 0.95:
            d = dfa(x[np.isfinite(x)])
            r = summary[(summary["metric"] == f"dfa_alpha_{name}")
                        & (summary["limb"] == limb)]
            if np.isfinite(d["alpha"]) and len(r):
                r = r.iloc[0]
                a.loglog(d["boxes"], d["fluct"], "o", color="#2a5d9f")
                lx = np.log10(d["boxes"])
                p = np.polyfit(lx, np.log10(d["fluct"]), 1)
                a.loglog(d["boxes"], 10 ** np.polyval(p, lx), color="#cc6633",
                         label=f"alpha {r['value']:.3f} "
                               f"[{r['ci_low']:.3f}, {r['ci_high']:.3f}]")
                a.set_title(f"DFA, {name}, {d['n']} strides ({limb})",
                            fontsize=10)
                a.legend(fontsize=8, frameon=False)
                break
    a.set_xlabel("box (strides)")
    a.set_ylabel("F(n)")

    # 2. alpha for every series
    a = ax[0, 1]
    d = summary[summary["metric"].str.startswith("dfa_alpha_")
                & (summary["limb"] == limb)].sort_values("value")
    if len(d):
        y = np.arange(len(d))
        err = np.vstack([d["value"] - d["ci_low"], d["ci_high"] - d["value"]])
        a.barh(y, d["value"], xerr=np.abs(err), color=[
            "#bbbbbb" if "metronome" in n else "#2a5d9f" for n in d["note"]],
            error_kw=dict(lw=0.8))
        a.set_yticks(y)
        a.set_yticklabels([m.replace("dfa_alpha_", "") for m in d["metric"]],
                          fontsize=7)
        a.axvline(0.5, color="#888", ls=":")
        a.axvline(1.0, color="#888", ls="--")
    a.set_title("alpha with 95% CI (grey: paced by the metronome)",
                fontsize=10)

    # 3. LDS divergence
    a = ax[1, 0]
    curves = extras.get("lds_curves") or {}
    if curves:
        tt = np.arange(1, LDS_WS * LDS_SAMPLES_PER_STRIDE + 1) \
            / LDS_SAMPLES_PER_STRIDE
        for name, c in curves.items():
            a.plot(tt, c, lw=1.5 if name == LDS_PRIMARY else 0.8,
                   label=name)
        a.set_xlabel("strides")
        a.set_ylabel("ln divergence")
        a.legend(fontsize=7, frameon=False)
    a.set_title("Local dynamic stability (mean over windows)", fontsize=10)

    # 4. multiscale entropy
    a = ax[1, 1]
    mse = extras.get("mse")
    if mse is not None:
        for c in mse.columns:
            a.plot(mse.index, mse[c], "o-" if c != "white_noise" else "s--",
                   ms=3, label=c)
        a.set_xlabel("scale (samples)")
        a.set_ylabel("sample entropy")
        a.legend(fontsize=7, frameon=False)
    a.set_title("Refined composite MSE, trunk acceleration", fontsize=10)
    for a in ax.flat:
        a.grid(alpha=0.15)
    fig.suptitle(t.stem, fontsize=12)
    fig.tight_layout()
    fig.savefig(path, dpi=110)
    plt.close(fig)


# %% --------------------------------------------------------------------------
#  birds-eye video of a typical stride
# -----------------------------------------------------------------------------

def typical_stride(t, limb="R"):
    """The steady stride closest to the median on stride time, step width
    and the two margins at contact, in IQR units."""
    s = t.steps[t.steps["steady"] & (t.steps["limb"] == limb)]
    cols = ["stride_s", "step_width_m", "mos_ml_contact", "mos_ap_contact"]
    s = s.dropna(subset=cols + ["next_hs", "to"])
    score = sum(np.abs(s[c] - s[c].median())
                / max(s[c].quantile(0.75) - s[c].quantile(0.25), 1e-9)
                for c in cols)
    return s.loc[score.idxmin()]


def mos_video(t, path, limb="R", fps=20, dpi=100):
    """Belt-frame birds-eye view: positions shifted forward by belt speed x
    time so the stance boot stays planted. Margins are differences at one
    instant, so the shift does not change them."""
    st = typical_stride(t, limb)
    rows = np.arange(int(np.ceil(st.hs)), int(np.floor(st.next_hs)) + 1)
    shift = t.belt_speed * (rows - rows[0]) / FS
    f2, l2 = t.forward, t.lateral

    def plane(xy, i):          # -> (along, across) in the belt frame
        return np.column_stack([xy @ f2 + shift[i], xy @ l2])

    hulls = {b: [] for b in LIMBS}
    for b in LIMBS:
        W = t.boots.world(b, rows)
        for i in range(len(rows)):
            p = plane(W[i, :, :2], i)
            hulls[b].append(p[ConvexHull(p).vertices])
    loaded = {b: np.zeros(len(rows), bool) for b in LIMBS}
    for x in t.steps.itertuples():
        if np.isfinite(x.to):
            loaded[x.limb] |= (rows >= x.hs) & (rows <= x.to)
    o = OTHER[limb]
    side = np.sign((t.ankle[limb][rows[0], :2] - t.ankle[o][rows[0], :2]) @ l2)
    length = np.linalg.norm(t.com[rows] - t.ankle[limb][rows], axis=1)
    w0 = np.sqrt(9.81 / length)
    com = np.array([plane(t.com[r:r + 1, :2], i)[0]
                    for i, r in enumerate(rows)])
    vel = np.column_stack([t.com_vel_belt[rows, :2] @ f2,
                           t.com_vel_belt[rows, :2] @ l2])
    xcom = com + vel / w0[:, None]
    ml = np.array([side * h[:, 1].max() if side > 0 else -h[:, 1].min()
                   for h in hulls[limb]]) - side * xcom[:, 1]
    ap = np.array([h[:, 0].max() for h in hulls[limb]]) - xcom[:, 0]
    stance = rows <= st.to

    allp = np.vstack(hulls["L"] + hulls["R"] + [com])
    lo, hi = allp.min(0) - 0.25, allp.max(0) + 0.25
    fig = plt.figure(figsize=(12, 7))
    top = fig.add_axes((0.06, 0.38, 0.9, 0.55))
    bot = fig.add_axes((0.06, 0.08, 0.9, 0.22))
    tt = (rows - rows[0]) / FS

    def draw(i):
        top.clear()
        bos = [hulls[b][i] for b in LIMBS if loaded[b][i]]
        if bos:
            p = np.vstack(bos)
            top.add_patch(Polygon(p[ConvexHull(p).vertices], closed=True,
                                  fc="#f2e5b8", ec="#8a6d1f", ls="--",
                                  alpha=0.6))
        for b in LIMBS:
            top.add_patch(Polygon(hulls[b][i], closed=True, lw=2,
                                  ec=COLOUR[b], alpha=0.8,
                                  fc=COLOUR[b] if loaded[b][i] else "none"))
        top.annotate("", xy=xcom[i], xytext=com[i],
                     arrowprops=dict(arrowstyle="-|>", color="k"))
        top.plot(*com[i], "ko", ms=9)
        top.plot(*xcom[i], "o", ms=9, mfc="none", mec="k", mew=2)
        top.set_xlim(lo[0], hi[0])
        top.set_ylim(lo[1], hi[1])
        top.set_aspect("equal")
        top.set_title(f"{SIDE[limb]} stride, {'stance' if stance[i] else 'swing'}"
                      f"  MoS ML {1000 * ml[i]:+.0f} mm  AP "
                      f"{1000 * ap[i]:+.0f} mm", fontsize=11)
        top.set_xlabel("walking direction, belt frame (m)")
        top.set_ylabel("across (m)")
        bot.clear()
        bot.plot(tt, 1000 * ml, color="#1b7a3d", label="MoS ML")
        bot.plot(tt, 1000 * ap, color="#7a4fb5", label="MoS AP")
        bot.axvline(tt[i], color="#888")
        bot.axhline(0, color="k", lw=0.8)
        bot.set_xlabel("s")
        bot.set_ylabel("mm")
        bot.legend(fontsize=8, frameon=False)

    from matplotlib.animation import FFMpegWriter, PillowWriter  # noqa
    try:
        writer, ext = FFMpegWriter(fps=fps), ".mp4"
        if not matplotlib.animation.writers.is_available("ffmpeg"):
            raise RuntimeError
    except Exception:
        writer, ext = PillowWriter(fps=fps), ".gif"
    out = str(path) + ext
    with writer.saving(fig, out, dpi):
        for i in range(len(rows)):
            draw(i)
            writer.grab_frame()
    plt.close(fig)
    return out


# %% ==========================================================================
# RUN: one trial from §5 to §19, then every trial
# =============================================================================

def read_participants(path=None):
    """participants.csv: participant, body_mass_kg (WITHOUT the load),
    leg_length_m, height_m, foot_length_m."""
    path = Path(path or PARTICIPANTS_FILE)
    if not path.exists():
        print(f"! {path} not found: no body mass, so the load is unknown")
        return {}
    p = pd.read_csv(path)
    return {r["participant"]: r.to_dict() for _, r in p.iterrows()}


def trial_stems():
    if TRIALS:
        return list(TRIALS)
    return sorted(p.name.replace("_merged_events.csv", "")
                  for p in Path(FOLDERS["events"]).glob("*_merged_events.csv"))


def steady_strides(stem):
    """Strides per foot after the warm-up, read off the event files (to cut
    every trial's series to the same length, §12.4)."""
    ev = Path(FOLDERS["events"])
    qa = ev / f"{stem}_event_qa.csv"
    e = pd.read_csv(qa if qa.exists() else ev / f"{stem}_merged_events.csv")
    rows = e["frame_float"] if "frame_float" in e else e["frame_100hz"]
    hs = e["event_type"] == HS
    start = rows.min() + WARMUP_S * FS
    return min(int(((e["support_limb"] == b) & hs & (rows >= start)).sum()) - 1
               for b in LIMBS)


def analyse_trial(stem, participant, series_length):
    """Everything for one trial, in the order of the Word document."""
    t = load_trial(stem, FOLDERS, participant)       # §5
    spatiotemporal(t)                                  # §6
    posture(t)                                         # §7
    kinetics(t)                                        # §8, §9
    margin_of_stability(t)                             # §10
    foot_clearance(t)                                  # §11
    print("  per-step metrics done; summarising (confidence intervals, DFA, "
          "entropy" + (", LDS" if RUN_LDS else "") + ")")
    table, extras = summarise(t, series_length)        # §12 - §18

    out = Path(OUTPUT_FOLDER)
    out.mkdir(parents=True, exist_ok=True)
    t.steps.to_csv(out / f"{stem}_steps.csv", index=False)
    table.to_csv(out / f"{stem}_summary.csv", index=False)
    extras["lds_windows"].to_csv(out / f"{stem}_lds_windows.csv", index=False)
    waves = {f"grf_{k}": np.asarray(v) for k, v in t.waves.items()}
    waves.update({f"angle_{k}": np.asarray(v)
                  for k, v in t.angle_waves.items()})
    waves.update({f"lds_{k}": v for k, v in extras["lds_curves"].items()})
    if extras["mse"] is not None:
        waves.update({f"mse_{c}": extras["mse"][c].to_numpy()
                      for c in extras["mse"]})
    np.savez_compressed(out / f"{stem}_waveforms.npz", **waves)
    if MAKE_FIGURES:                                   # §19
        qa_figure(t, out / f"{stem}_qa.png")
        variability_figure(t, table, extras, series_length,
                           out / f"{stem}_variability.png", LIMB_FOR_TRUNK)
    if MAKE_VIDEO:
        print(f"  video: {mos_video(t, out / f'{stem}_mos_stride')}")
    print(f"  saved to {out}")
    return t, table


def main(stems=None):
    stems = stems or trial_stems()
    people = read_participants()
    n = SERIES_LENGTH
    if n == "shortest":          # §12.4: every trial at the same N
        counts = {s: steady_strides(s) for s in stems}
        n = min(counts.values())
        print(f"series cut to {n} strides, the shortest trial's "
              f"({min(counts, key=counts.get)})")
    done = 0
    for stem in stems:
        try:
            analyse_trial(stem, people.get(stem.split("_")[0], {}), n)
            done += 1
        except Exception as exc:             # keep the batch going
            print(f"\n{stem}: FAILED -- {exc}")
            traceback.print_exc()
    print(f"\n{done} of {len(stems)} trials analysed")
    return build_tables(OUTPUT_FOLDER)       # §19: every trial, one row each


if __name__ == "__main__":
    if sys.argv[1:] == ["--tables"]:
        build_tables(OUTPUT_FOLDER)
    else:
        main(sys.argv[1:] or None)
