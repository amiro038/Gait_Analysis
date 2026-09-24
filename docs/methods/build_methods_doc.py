# -*- coding: utf-8 -*-
"""
Build docs/Gait_Analysis_Methods.docx from the markdown in docs/methods/.

    python docs/methods/build_methods_doc.py

The text is written in markdown with LaTeX equations (native Word equations
after conversion). Tables that must agree with the code are generated FROM
the code, so they cannot drift apart:

    {{ranges: col, col, ...}}   healthy ranges of those results columns
                                (gait_analysis.REFERENCE_RANGES)
    {{allranges}}               every healthy range (Appendix B)
    {{params: file}}            every setting of that script, with its comment
    {{codemap}}                 every section of gait_analysis.py and its
                                functions

Needs pypandoc_binary and python-docx (pip install pypandoc_binary
python-docx). Figures come from docs/make_methods_figures.py.
"""

import re
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
import gait_analysis as ga                                  # noqa: E402

PARTS = sorted(HERE.glob("[0-9][0-9]_*.md"))
OUT = ROOT / "docs" / "Gait_Analysis_Methods.docx"


def cell(x):
    return str(x).replace("|", "\\|").replace("\n", " ")


def fmt(v):
    return f"{v:g}"


def ranges_table(cols):
    # the dashes set the relative column widths (pandoc)
    head = ("| Results column | What it is | Healthy range | Population | Source | Checked |\n"
            "|:" + "-" * 24 + "|:" + "-" * 21 + "|:" + "-" * 13 + "|:" + "-" * 19
            + "|:" + "-" * 14 + "|:" + "-" * 9 + "|\n")
    rows = []
    for c in cols:
        ref = ga.reference_for(c)
        if ref is None:
            raise KeyError(f"no healthy range for {c}")
        lo, hi, pop, src, conf = ref
        _, desc, unit, _ = ga.describe(c[:-3] if c.endswith(("_sd", "_cv")) else c)
        if c.endswith("_cv"):
            desc, unit = f"{desc}: CV over steps", "%"
        elif c.endswith("_sd"):
            desc = f"{desc}: SD over steps"
        unit = "" if unit in ("-", "") else f" {unit}"
        rows.append(f"| `{c}` | {cell(desc)} | {fmt(lo)} to {fmt(hi)}{unit} | {cell(pop)} "
                    f"| {cell(src)} | {'yes' if conf == 'confirmed' else 'verify'} |")
    return head + "\n".join(rows) + "\n"


def all_ranges():
    return ranges_table(list(ga.REFERENCE_RANGES))


SETTING_NOTES = {       # settings whose comment in the script is shared or missing
    "GRF_RATE": "force-plate sampling rate, Hz (§2.2)",
    "KINEMATIC_RATE": "Theia sampling rate, Hz (§2.1)",
    "COP_FRAME": "frame of the exported CoP: \"plate\" = each belt's own frame, origin at its centre (§5.4)",
    "EDGE_HOLD": "samples the raw force must stay across the threshold at an edge (§4.1)",
    "TOE_SIGN": "sign of Theia's toe angle for extension (§3.3)",
    "SHARED_MIN_FRAMES": "frames with the other boot on this belt that make a contact \"shared\" (§4.3)",
    "STRADDLE_MIN_FRAMES": "frames with this boot also on the other belt that make it a \"straddle\" (§4.3)",
    "COP_MIN_FORCE_N": "the CoP is used only above this force (§4.3, §5.4)",
    "ZENI_WINDOW_S": "window either side for the prominence test (§4.4)",
    "CI_LEVEL": "confidence level of every interval (§12)",
    "ENT_R_SWEEP": "tolerances also reported, to show the dependence on r (§15.1)",
    "RUN_LDS": "compute §18 (the slowest step)",
    "MAKE_FIGURES": "write the per-trial figures (§19.3)",
}


def params_table(script):
    """Every NAME = value line of a script's settings block, with its comment:
    the one on the same line, or else the block of comment lines just above
    it (given to the first setting under the block only)."""
    lines = (ROOT / script).read_text(encoding="utf-8").splitlines()
    if script == "gait_analysis.py":
        a = next(i for i, l in enumerate(lines) if l.startswith("# §0"))
        b = next(i for i, l in enumerate(lines) if l.startswith("LIMBS, OTHER"))
    else:
        a = next(i for i, l in enumerate(lines) if l.startswith("#  CONFIG"))
        b = next(i for i, l in enumerate(lines) if l.startswith("LIMBS = "))
    rows, last, block = [], None, []
    for l in lines[a + 1:b]:
        if re.match(r"# (---|===|%%)", l):
            block, last = [], None
            continue
        if l.startswith("#"):
            block.append(l.lstrip("#").strip())
            last = None
            continue
        m = re.match(r"^([A-Z][A-Z0-9_]+)\s*=\s*(.+?)(?:\s+#\s*(.*))?$", l)
        if m and not l.startswith(("FOLDERS", "LDS_STATE_SPACES", "TRUNK_SOURCES")):
            name, value, comment = m.groups()
            if not comment and block:
                comment = " ".join(block)
            block = []
            if ("Path(" in value or "DATA_FOLDER" in value or "HERE /" in value
                    or value.startswith(("dict(", "["))):
                last = None
                continue
            comment = SETTING_NOTES.get(name, comment)
            last = [name, value.rstrip(",").strip(), (comment or "").strip()]
            rows.append(last)
        elif last is not None and re.match(r"^\s{10,}#\s*(.*)$", l):
            last[2] += " " + re.match(r"^\s{10,}#\s*(.*)$", l).group(1).strip()
        elif not l.strip():
            block, last = [], None
        elif not l.startswith(" "):
            last = None
    out = ("| Setting | Value | What it does (section) |\n|:" + "-" * 26 + "|:" + "-" * 22
           + "|:" + "-" * 52 + "|\n")
    for name, value, comment in rows:
        out += f"| `{name}` | `{cell(value.replace('`', ''))}` | {cell(comment)} |\n"
    return out


def codemap():
    lines = (ROOT / "gait_analysis.py").read_text(encoding="utf-8").splitlines()
    sections, cur = [], None
    for k, l in enumerate(lines):
        header = k > 0 and lines[k - 1].startswith("# %% ====")
        if header:
            m = re.match(r"^# (§[0-9.]+)\s+(.+?)\s*(?:\(Word document.*\))?$", l)
            if m:
                cur = [m.group(1), m.group(2).strip().rstrip(","), []]
            elif l.startswith("# RUN"):
                cur = ["RUN", "one trial, then every trial", []]
            elif l.startswith("# SMALL TOOLS"):
                cur = ["tools", "small tools used throughout", []]
            else:
                continue
            sections.append(cur)
            continue
        f = re.match(r"^(?:def|class) (\w+)", l)
        if f and cur is not None:
            cur[2].append(f.group(1))

    def sentence(t):                       # "MARGIN OF STABILITY" -> "Margin of stability"
        t = t[0] + t[1:].lower()
        for word in ("CoM", "DFA", "GEM"):
            t = re.sub(rf"\b{word}\b", word, t, flags=re.I)
        return t

    out = ("| Section | What | Functions |\n|:" + "-" * 10 + "|:" + "-" * 30 + "|:"
           + "-" * 60 + "|\n")
    for sec, title, funcs in sections:
        if sec.startswith("§0"):
            funcs = ["(settings only)"]
        out += f"| {sec} | {cell(sentence(title))} | {', '.join(f'`{x}`' for x in funcs) or '-'} |\n"
    return out


def expand(text):
    text = re.sub(r"\{\{ranges:\s*([^}]+)\}\}",
                  lambda m: ranges_table([c.strip() for c in m.group(1).split(",")]), text)
    text = text.replace("{{allranges}}", all_ranges())
    text = re.sub(r"\{\{params:\s*([^}]+)\}\}", lambda m: params_table(m.group(1).strip()), text)
    text = text.replace("{{codemap}}", codemap())
    return text


# --- the table of contents, filled in ------------------------------------------
# pandoc writes Word's contents field empty (Word fills it when the document
# is opened). To show it filled in every viewer, the entries are written into
# the field here, with page numbers read from a LibreOffice rendering (Calibri
# and its metric twin Carlito paginate alike). The field stays marked "dirty",
# so Word offers to refresh it on opening; right-click -> Update Field also does.

def xml_escape(x):
    return x.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def headings(xml):
    """(level, text, bookmark) of every Heading 1 / Heading 2, in order. pandoc
    puts each heading's bookmark just before its paragraph."""
    out = []
    for m in re.finditer(r'<w:p><w:pPr><w:pStyle w:val="Heading([12])" ?/>.*?</w:p>', xml, flags=re.S):
        marks = re.findall(r'<w:bookmarkStart w:id="\d+" w:name="([^"]+)" ?/>\s*$', xml[max(m.start() - 400, 0):m.start()])
        if not marks:
            continue
        text = "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", m.group(0)))
        text = text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
        out.append((int(m.group(1)), text, marks[-1]))
    return out


def toc_xml(entries, pages):
    paras = []
    for i, (level, text, mark) in enumerate(entries):
        link = (f'<w:hyperlink w:anchor="{mark}" w:history="1"><w:r><w:t xml:space="preserve">'
                f'{xml_escape(text)}</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>{pages[i]}</w:t></w:r>'
                f'</w:hyperlink>')
        head = ('<w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r><w:r><w:instrText '
                'xml:space="preserve">TOC \\o &quot;1-2&quot; \\h \\z \\u</w:instrText></w:r>'
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r>') if i == 0 else ""
        tail = '<w:r><w:fldChar w:fldCharType="end"/></w:r>' if i == len(entries) - 1 else ""
        paras.append(f'<w:p><w:pPr><w:pStyle w:val="TOC{level}"/></w:pPr>{head}{link}{tail}</w:p>')
    return "".join(paras)


def rewrite(docx, fn):
    import zipfile
    tmp = docx.with_suffix(".tmp")
    with zipfile.ZipFile(docx) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = fn(data.decode("utf-8")).encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(docx)


def fill_toc(docx, pages=None):
    def fn(xml):
        entries = headings(xml)
        nums = pages or ["00"] * len(entries)
        field = re.search(r'<w:p><w:r><w:fldChar w:fldCharType="begin".*?</w:p>', xml, flags=re.S)
        if field is None:                        # already filled: replace the entries
            field = re.search(r'<w:p><w:pPr><w:pStyle w:val="TOC1"/>.*<w:fldChar w:fldCharType="end"/></w:r></w:p>',
                              xml, flags=re.S)
        return xml[:field.start()] + toc_xml(entries, nums) + xml[field.end():]
    rewrite(docx, fn)


def render(docx, folder):
    """PDF of the document via LibreOffice, or None if it is not installed."""
    import shutil
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(folder), str(docx)],
                   check=True, capture_output=True)
    pdf = Path(folder) / (docx.stem + ".pdf")
    return pdf if pdf.exists() else None


def heading_pages(pdf, entries, body_starts_with):
    import pymupdf
    doc = pymupdf.open(str(pdf))
    text = [" ".join(p.get_text().split()) for p in doc]
    first = next(i for i, t in enumerate(text) if body_starts_with in t)
    pages, k = [], first
    for _, title, _ in entries:
        key = " ".join(title.split())[:45]
        while k < len(text) and key not in text[k]:
            k += 1
        if k == len(text):
            raise RuntimeError(f"heading not found in the rendering: {title}")
        pages.append(str(k + 1))
    return pages


def small_tables(docx):
    """9-pt text in every table: the paragraph style inside cells (Compact)
    is also used by lists, so the size is set on the table text itself."""
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    d = Document(str(docx))
    for table in d.tables:
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            for c in row.cells:
                for para in c.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
    d.save(str(docx))


def main():
    import pypandoc
    sys.path.insert(0, str(HERE))
    import make_reference_docx
    ref = make_reference_docx.main(HERE / "reference.docx")
    md = "\n\n".join(expand(p.read_text(encoding="utf-8")) for p in PARTS)
    src = HERE / "_combined.md"
    src.write_text(md, encoding="utf-8")
    subprocess.run([pypandoc.get_pandoc_path(), str(src), "-o", str(OUT),
                    "--from", "markdown+pipe_tables+tex_math_dollars+implicit_figures",
                    "--reference-doc", str(ref), "--toc", "--toc-depth=2",
                    "--resource-path", str(ROOT / "docs")], check=True)
    small_tables(OUT)
    fill_toc(OUT)                                 # placeholder numbers first ...
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        pdf = render(OUT, tmp)                    # ... so the pagination is final
        if pdf is None:
            print("  (no LibreOffice: the contents are filled in when Word opens the file)")
        else:
            import zipfile
            xml = zipfile.ZipFile(OUT).read("word/document.xml").decode("utf-8")
            pages = heading_pages(pdf, headings(xml), "This document and the script")
            fill_toc(OUT, pages)
    print(f"{OUT}  ({len(md.split())} words)")


if __name__ == "__main__":
    main()
