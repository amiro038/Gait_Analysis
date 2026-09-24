"""
The results as a table you can read: one row per trial, one column per
metric.

    all_trials_metrics.csv      the values
    all_trials_metrics_ci.csv   the 95% confidence limits of the same columns
                                (<column>_lo, <column>_hi)
    metric_dictionary.csv       every column: what it is, its unit, which
                                limb and statistic, and the section of
                                METHODS.md that defines it
    all_trials_metrics.xlsx     the same in a workbook: a "Key metrics" sheet,
                                one sheet per family of metrics, the
                                confidence limits and the dictionary, with a
                                readable header above the column names

They are rebuilt from every {trial}_summary.csv in the output folder, so a
trial analysed yesterday and one analysed today end up in the same table.

HOW THE COLUMNS ARE NAMED
    step_width_m            mean over all steady steps, both feet
    step_width_m_sd         SD over the same steps
    step_width_m_cv         CV (%)
    step_width_m_L / _R     mean of the left / right foot's steps
    dfa_alpha_step_width_m_L    a per-limb value (DFA, entropy, GEM, foot
                                placement are computed on each foot's own
                                series)
Everything else is one value per trial. Units are the end of the name:
_s, _m, _mm, _pct, _deg, _bw (body weight), _tw (total weight, body +
load), _j_kg, _ms (m/s), _spm (steps/min).
"""

from pathlib import Path

import numpy as np
import pandas as pd

# name: (family, description, unit, METHODS.md section)
CATALOG = {
    # --- the trial ------------------------------------------------------------
    "belt_speed_ms": ("Trial", "Belt speed, measured from the stance heels", "m/s", "5.2"),
    "system_mass_kg": ("Trial", "Body + load, weighed in the quiet standing", "kg", "5.4"),
    "load_kg": ("Trial", "Carried load: system mass minus body mass", "kg", "5.4"),
    "load_offset_y_m": ("Trial", "Load CoM forward of Trunk_Position (negative = behind)", "m", "5.4"),
    "load_offset_x_m": ("Trial", "Load CoM to the right of Trunk_Position (0 when centred)", "m", "5.4"),
    "load_offset_z_m": ("Trial", "Load CoM above Trunk_Position (assumed)", "m", "5.4"),
    "quiet_standing_cv": ("Trial", "CV of the total vertical force in the quiet standing window", "-", "5.4"),
    "steps_steady": ("Trial", "Steps after the warm-up", "steps", "5.1"),
    "events_trusted_grf_pct": ("Trial", "Events from trusted GRF", "%", "4"),
    "events_zeni_pct": ("Trial", "Events from Zeni (kinematic fallback)", "%", "4"),
    "events_interpolated_pct": ("Trial", "Events interpolated (last resort)", "%", "4"),
    "stances_kinetics_pct": ("Trial", "Steady stances trusted for kinetics", "%", "8"),
    "swings_without_mfc_pct": ("Trial", "Swings with no valid MFC event (non-MTC)", "%", "11"),
    "mfc_below_belt_pct": ("Trial", "Swings whose MFC is below the belt surface (pose error)", "%", "11"),
    "handrail_steps_pct": ("Trial", "Steady steps with handrail contact", "%", "5.5"),
    "grf_vs_marker_acc_r_X": ("Trial", "GRF vs marker CoM acceleration, Theia X (0.5-5 Hz)", "r", "5.5"),
    "grf_vs_marker_acc_r_Y": ("Trial", "GRF vs marker CoM acceleration, Theia Y (0.5-5 Hz)", "r", "5.5"),
    "grf_vs_marker_acc_r_Z": ("Trial", "GRF vs marker CoM acceleration, vertical (0.5-5 Hz)", "r", "5.5"),
    # --- spatiotemporal --------------------------------------------------------
    "stride_s": ("Spatiotemporal", "Stride time", "s", "6"),
    "stance_s": ("Spatiotemporal", "Stance time", "s", "6"),
    "swing_s": ("Spatiotemporal", "Swing time", "s", "6"),
    "step_s": ("Spatiotemporal", "Step time", "s", "6"),
    "cadence_spm": ("Spatiotemporal", "Cadence", "steps/min", "6"),
    "stance_pct": ("Spatiotemporal", "Stance, % of stride", "%", "6"),
    "double_support_pct": ("Spatiotemporal", "Double support, % of stride", "%", "6"),
    "single_support_pct": ("Spatiotemporal", "Single support, % of stride", "%", "6"),
    "step_length_m": ("Spatiotemporal", "Step length, heel to heel along the belt", "m", "6"),
    "step_width_m": ("Spatiotemporal", "Step width, heel to heel across the belt", "m", "6"),
    "stride_length_m": ("Spatiotemporal", "Stride length over the belt", "m", "6"),
    "stride_speed_ms": ("Spatiotemporal", "Stride speed over the belt", "m/s", "6"),
    "walk_ratio": ("Spatiotemporal", "Walk ratio: step length / cadence", "m/(steps/min)", "12.3"),
    # --- posture -----------------------------------------------------------------
    "trunk_lean_mean_deg": ("Posture", "Trunk (thorax) lean, mean over the stride", "deg", "7"),
    "trunk_lean_rom_deg": ("Posture", "Trunk lean, range within the stride", "deg", "7"),
    "pelvis_tilt_mean_deg": ("Posture", "Pelvis tilt, mean over the stride", "deg", "7"),
    "hip_rom_deg": ("Posture", "Hip sagittal range of motion", "deg", "7"),
    "knee_rom_deg": ("Posture", "Knee sagittal range of motion", "deg", "7"),
    "ankle_rom_deg": ("Posture", "Ankle sagittal range of motion", "deg", "7"),
    # --- kinetics ------------------------------------------------------------------
    "f1_bw": ("Kinetics", "Vertical GRF first peak (weight acceptance)", "BW", "8"),
    "trough_bw": ("Kinetics", "Vertical GRF mid-stance trough", "BW", "8"),
    "f2_bw": ("Kinetics", "Vertical GRF second peak (push-off)", "BW", "8"),
    "f1_tw": ("Kinetics", "Vertical GRF first peak, per total weight", "TW", "8"),
    "f2_tw": ("Kinetics", "Vertical GRF second peak, per total weight", "TW", "8"),
    "loading_rate_bw_per_s": ("Kinetics", "Loading rate, 20-80% of F1", "BW/s", "8"),
    "peak_braking_bw": ("Kinetics", "Peak braking force", "BW", "8"),
    "peak_propulsive_bw": ("Kinetics", "Peak propulsive force", "BW", "8"),
    "braking_impulse_bw_s": ("Kinetics", "Braking impulse (negative)", "BW s", "8"),
    "propulsive_impulse_bw_s": ("Kinetics", "Propulsive impulse", "BW s", "8"),
    "vertical_impulse_bw_s": ("Kinetics", "Vertical impulse", "BW s", "8"),
    "braking_impulse_tw_s": ("Kinetics", "Braking impulse, per total weight", "TW s", "8"),
    "propulsive_impulse_tw_s": ("Kinetics", "Propulsive impulse, per total weight", "TW s", "8"),
    "cop_ap_range_mm": ("Kinetics", "CoP excursion along the belt", "mm", "8"),
    "cop_ml_range_mm": ("Kinetics", "CoP excursion across the belt", "mm", "8"),
    "free_moment_peak_nm": ("Kinetics", "Peak |free moment|", "N m", "8"),
    # --- CoM work ----------------------------------------------------------------
    "collision_j_kg": ("CoM work", "Collision: negative work, heel strike to other toe-off", "J/kg", "9"),
    "rebound_j_kg": ("CoM work", "Rebound: positive work in single support", "J/kg", "9"),
    "preload_j_kg": ("CoM work", "Preload: negative work in single support", "J/kg", "9"),
    "push_off_j_kg": ("CoM work", "Push-off: positive work, other heel strike to toe-off", "J/kg", "9"),
    "com_work_pos_j_kg": ("CoM work", "Positive work over the stance", "J/kg", "9"),
    "com_work_neg_j_kg": ("CoM work", "Negative work over the stance", "J/kg", "9"),
    # --- stability -----------------------------------------------------------------
    "mos_ml_contact": ("Stability", "ML margin of stability at heel strike", "m", "10"),
    "mos_ml_min": ("Stability", "ML margin of stability, minimum in single support", "m", "10"),
    "mos_ap_contact": ("Stability", "AP margin of stability at heel strike", "m", "10"),
    "mos_ap_min": ("Stability", "AP margin of stability, minimum in single support", "m", "10"),
    "mfc_m": ("Stability", "Minimum foot clearance above the belt", "m", "11"),
    "moi_peak_mm": ("Stability", "Margin of instability, peak in swing", "mm", "11"),
    "tri_s": ("Stability", "Trip risk integral", "s", "11"),
    # --- trunk ---------------------------------------------------------------------
    "harmonic_ratio_hr_AP": ("Trunk", "Harmonic ratio AP (median over strides)", "-", "15.1"),
    "harmonic_ratio_hr_VT": ("Trunk", "Harmonic ratio vertical (median over strides)", "-", "15.1"),
    "harmonic_ratio_hr_ML": ("Trunk", "Harmonic ratio ML (median over strides)", "-", "15.1"),
    "harmonic_ratio_ihr_AP": ("Trunk", "Improved harmonic ratio AP (mean)", "%", "15.1"),
    "harmonic_ratio_ihr_VT": ("Trunk", "Improved harmonic ratio vertical (mean)", "%", "15.1"),
    "harmonic_ratio_ihr_ML": ("Trunk", "Improved harmonic ratio ML (mean)", "%", "15.1"),
    "step_regularity_AP": ("Trunk", "Step regularity AP", "-", "15.2"),
    "step_regularity_VT": ("Trunk", "Step regularity vertical", "-", "15.2"),
    "step_regularity_ML": ("Trunk", "Step regularity ML", "-", "15.2"),
    "stride_regularity_AP": ("Trunk", "Stride regularity AP", "-", "15.2"),
    "stride_regularity_VT": ("Trunk", "Stride regularity vertical", "-", "15.2"),
    "stride_regularity_ML": ("Trunk", "Stride regularity ML", "-", "15.2"),
    "symmetry_AP": ("Trunk", "Step / stride regularity, AP", "-", "15.2"),
    "symmetry_VT": ("Trunk", "Step / stride regularity, vertical", "-", "15.2"),
    "symmetry_ML": ("Trunk", "Step / stride regularity, ML", "-", "15.2"),
    "trunk_acc_rms_AP": ("Trunk", "Trunk acceleration RMS, AP", "m/s2", "15"),
    "trunk_acc_rms_VT": ("Trunk", "Trunk acceleration RMS, vertical", "m/s2", "15"),
    "trunk_acc_rms_ML": ("Trunk", "Trunk acceleration RMS, ML", "m/s2", "15"),
    "mse_area_AP": ("Trunk", "Multiscale entropy, area over scales 1-30, AP", "-", "14.2"),
    "mse_area_VT": ("Trunk", "Multiscale entropy, area, vertical", "-", "14.2"),
    "mse_area_ML": ("Trunk", "Multiscale entropy, area, ML", "-", "14.2"),
    # --- control -------------------------------------------------------------------
    "foot_placement_r2": ("Control", "ML foot placement explained by CoM state (R2)", "-", "16.2"),
    "gem_perpendicular_sd_pct": ("Control", "GEM goal-relevant SD", "% stride", "16.1"),
    "gem_parallel_sd_pct": ("Control", "GEM goal-equivalent SD", "% stride", "16.1"),
    "gem_perpendicular_lag1": ("Control", "GEM goal-relevant lag-1 autocorrelation", "-", "16.1"),
    "gem_parallel_lag1": ("Control", "GEM goal-equivalent lag-1 autocorrelation", "-", "16.1"),
    "gem_perpendicular_dfa_alpha": ("Control", "GEM goal-relevant DFA alpha", "-", "16.1"),
    "gem_parallel_dfa_alpha": ("Control", "GEM goal-equivalent DFA alpha", "-", "16.1"),
}
# the order of the families in the tables; DFA, entropy, symmetry and LDS
# columns are named from the series they are computed on (see describe)
FAMILIES = ["Trial", "Spatiotemporal", "Stability", "Kinetics", "CoM work",
            "Posture", "Symmetry", "DFA", "Entropy", "Trunk", "Control",
            "LDS"]
KEY = ["load_kg", "system_mass_kg", "belt_speed_ms", "steps_steady",
       "events_trusted_grf_pct",
       "stride_s", "cadence_spm", "stance_pct", "double_support_pct",
       "step_length_m", "step_width_m", "step_width_m_cv",
       "stride_length_m_cv",
       "mos_ml_contact", "mos_ml_min", "mos_ap_min",
       "mfc_m", "mfc_m_sd", "tri_s",
       "f1_bw", "f2_bw", "loading_rate_bw_per_s",
       "braking_impulse_bw_s", "propulsive_impulse_bw_s",
       "collision_j_kg", "push_off_j_kg",
       "trunk_lean_mean_deg", "knee_rom_deg",
       "dfa_alpha_stride_s_L", "dfa_alpha_stride_s_R",
       "dfa_alpha_step_width_m_L", "dfa_alpha_step_width_m_R",
       "foot_placement_r2_L", "foot_placement_r2_R",
       "harmonic_ratio_ihr_AP", "lds_lambda_S_trunkVel_AP"]
ID = ["participant", "condition", "trial"]


def describe(metric):
    """(family, description, unit, section) for any metric name in the long
    table, including the families built from a series name."""
    if metric in CATALOG:
        return CATALOG[metric]
    for prefix, family, what, section in (
            ("dfa_alpha_", "DFA", "DFA alpha of", "13"),
            ("sample_entropy_", "Entropy", "Sample entropy (m 2, r 0.2 SD) of",
             "14.1"),
            ("symmetry_angle_", "Symmetry", "Symmetry angle (0 = symmetric, "
             "+ = right larger) of", "12.2")):
        if metric.startswith(prefix):
            base = metric[len(prefix):]
            label = CATALOG.get(base, ("", base))[1].lower()
            return (family, f"{what} {label}", "%" if family == "Symmetry"
                    else "-", section)
    if metric.startswith("lds_lambda_"):
        tag, space = metric[len("lds_lambda_"):].split("_", 1)
        span = "short-term (0-0.5 stride)" if tag == "S" else \
            "long-term (4-10 strides)"
        return ("LDS", f"Lyapunov exponent, {span}, state space {space}",
                "1/stride", "17")
    return ("Other", metric, "", "")


def column(metric, limb, statistic):
    """Wide column name for one row of the long table (None = not shown)."""
    if metric.startswith("harmonic_ratio_hr_"):     # unbounded: the median
        return metric if statistic == "median" else None
    if metric.startswith("harmonic_ratio_ihr_"):    # bounded: the mean
        return metric if statistic == "mean" else None
    if metric.startswith("lds_"):                   # one limb's strides
        return metric
    if statistic in ("sd", "cv"):
        return f"{metric}_{statistic}" if limb == "both" else None
    if statistic in ("mean", "value", "pct"):
        return metric if limb == "both" else f"{metric}_{limb}"
    return None


def wide(long):
    """Long summary rows -> (values, CI limits, dictionary)."""
    long = long.copy()
    long["column"] = [column(m, b, s) for m, b, s in
                      zip(long["metric"], long["limb"], long["statistic"])]
    long = long.dropna(subset=["column"])
    order = {f: i for i, f in enumerate(FAMILIES + ["Other"])}
    meta = {}
    for m, b, s, c in zip(long["metric"], long["limb"], long["statistic"],
                          long["column"]):
        fam, desc, unit, sec = describe(m)
        foot = {"L": "left foot", "R": "right foot", "both": "both feet"}[b]
        if s in ("sd", "cv"):
            desc = f"{desc}, {'SD' if s == 'sd' else 'CV'} over steps ({foot})"
            unit = "%" if s == "cv" else unit
        elif s == "mean" and fam not in ("Trunk",):
            desc = f"{desc}, mean over steps ({foot})"
        elif b in ("L", "R") and c.endswith(f"_{b}"):
            desc = f"{desc}, {foot}"
        meta[c] = dict(column=c, family=fam, description=desc, unit=unit,
                       limb=b, statistic=s, methods_section=sec, metric=m)
    # by family, then each metric where it first appears, then: mean of
    # both feet, SD, CV, left, right
    first = list(dict.fromkeys(long["metric"]))
    variant = {"both": 0, "L": 3, "R": 4}
    cols = sorted(meta, key=lambda c: (
        order[meta[c]["family"]], first.index(meta[c]["metric"]),
        {"sd": 1, "cv": 2}.get(meta[c]["statistic"],
                               variant[meta[c]["limb"]])))
    idx = ["participant", "condition", "trial"]
    values = long.pivot_table(index=idx, columns="column", values="value",
                              aggfunc="first")
    lo = long.pivot_table(index=idx, columns="column", values="ci_low",
                          aggfunc="first")
    hi = long.pivot_table(index=idx, columns="column", values="ci_high",
                          aggfunc="first")
    values = values.reindex(columns=cols).reset_index()
    values.columns.name = None
    ci_cols = [c for c in cols if c in lo and lo[c].notna().any()]
    ci = pd.concat([lo[ci_cols].add_suffix("_lo"), hi[ci_cols].add_suffix("_hi")],
                   axis=1)
    ci = ci[[f"{c}_{e}" for c in ci_cols for e in ("lo", "hi")]].reset_index()
    ci.columns.name = None
    dictionary = pd.DataFrame([meta[c] for c in cols]).drop(columns="metric")
    return values, ci, dictionary


def build(folder):
    """Rebuild the wide tables from every {trial}_summary.csv in folder."""
    folder = Path(folder)
    files = sorted(folder.glob("*_summary.csv"))
    files = [f for f in files if not f.name.startswith("all_trials")]
    if not files:
        return None
    long = pd.concat([pd.read_csv(f) for f in files], ignore_index=True)
    long.to_csv(folder / "all_trials_summary.csv", index=False)
    values, ci, dictionary = wide(long)
    values.to_csv(folder / "all_trials_metrics.csv", index=False)
    ci.to_csv(folder / "all_trials_metrics_ci.csv", index=False)
    dictionary.to_csv(folder / "metric_dictionary.csv", index=False)
    xlsx = folder / "all_trials_metrics.xlsx"
    try:
        write_workbook(xlsx, values, ci, dictionary)
    except ImportError:
        print("  (install openpyxl for the .xlsx workbook; the CSVs are "
              "written)")
        xlsx = None
    print(f"\n{len(values)} trials x {values.shape[1] - 3} metrics -> "
          f"{folder / 'all_trials_metrics.csv'}"
          + (f"\n  and {xlsx.name}" if xlsx else ""))
    return values, ci, dictionary


def write_workbook(path, values, ci, dictionary):
    """One sheet per family, the key metrics first, a readable header row
    (description and unit) above the column names, trials frozen on the
    left."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    info = dictionary.set_index("column")
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, cols, frame=values, labels=None):
        ws = wb.create_sheet(name[:31])
        cols = ID + [c for c in cols if c in frame]
        for j, c in enumerate(cols, 1):
            if labels is not None:
                label = labels.get(c, c)
            elif c in info.index:
                label = f"{info.at[c, 'description']} ({info.at[c, 'unit']})"
            else:
                label = c.capitalize()
            top = ws.cell(row=1, column=j, value=label)
            top.font = Font(bold=True)
            top.alignment = Alignment(wrap_text=True, vertical="top")
            top.fill = PatternFill("solid", fgColor="DDE6F0")
            low = ws.cell(row=2, column=j, value=c)
            low.font = Font(italic=True, color="666666", size=9)
            for i, v in enumerate(frame[c].tolist(), 3):
                if isinstance(v, float) and not np.isfinite(v):
                    v = None
                cell = ws.cell(row=i, column=j, value=v)
                if isinstance(v, float):
                    cell.number_format = "0.000" if abs(v) < 100 else "0.0"
            ws.column_dimensions[get_column_letter(j)].width = \
                34 if c == "trial" else 16
        ws.row_dimensions[1].height = 60
        ws.freeze_panes = "D3"

    sheet("Key metrics", KEY)
    for fam in FAMILIES + ["Other"]:
        cols = dictionary.loc[dictionary["family"] == fam, "column"].tolist()
        if cols:
            sheet(fam, cols)
    ci_labels = {}
    for c in ci.columns:
        base = c[:-3]
        if base in info.index:
            ci_labels[c] = (f"{info.at[base, 'description']}, 95% CI "
                            f"{'lower' if c.endswith('_lo') else 'upper'}")
    sheet("95% CI", [c for c in ci.columns if c not in ID], ci, ci_labels)
    ws = wb.create_sheet("Dictionary")
    head = list(dictionary.columns)
    for j, h in enumerate(head, 1):
        ws.cell(row=1, column=j, value=h).font = Font(bold=True)
    for i, r in enumerate(dictionary.itertuples(index=False), 2):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
    for j, w in enumerate((34, 14, 70, 12, 6, 10, 10), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"
    wb.save(path)
